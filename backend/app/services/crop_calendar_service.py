"""
Crop Calendar & Agricultural Catalog Service

Provides:
1. Dynamic loading and normalization of all 181 agri-products across 7 Excel catalog files
   (fertilizers, insecticides, fungicides, herbicides, seeds, bio-inputs, adjuvants).
2. Comprehensive Crop Lifecycle Engine mapping sowing dates to agronomic growth stages,
   days to harvest, stage-specific fertilizer/input requirements, and urgency timelines.
3. Matching engine to calculate stage-aware stocking demand for retail agricultural shops.
"""

import os
import re
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Any, Optional, Tuple

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# 1. PATH RESOLUTION & EXCEL CATALOG LOADER
# ---------------------------------------------------------------------------

def _resolve_data_dir() -> Optional[str]:
    """Finds the directory containing the 7 agricultural Excel files."""
    candidates = [
        r"C:\Users\vinay\OneDrive\Desktop\agri\data",
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "data")),
        os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "data")),
        os.path.abspath(os.path.join(os.getcwd(), "data")),
        os.path.abspath(os.path.join(os.getcwd(), "..", "data")),
    ]
    for p in candidates:
        if os.path.exists(p) and os.path.isdir(p):
            return p
    return None


DATA_DIR = _resolve_data_dir()

# Global in-memory cache of normalized products
_CATALOG_CACHE: Optional[List[Dict[str, Any]]] = None
_CATALOG_BY_ID: Dict[str, Dict[str, Any]] = {}
_CATALOG_BY_CATEGORY: Dict[str, List[Dict[str, Any]]] = {}


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    return "" if s.lower() in ("none", "nan", "null") else s


def _extract_crops_list(text: str) -> List[str]:
    """Parses free-text target crops into standard crop names."""
    if not text:
        return []
    clean = text.replace("–", "-").replace("—", "-")
    # If "All crops", record that
    crops = []
    is_all = "all crop" in clean.lower() or "universal" in clean.lower()
    
    known_crops = [
        ("Rice", "Paddy (Rice)"), ("Paddy", "Paddy (Rice)"), ("Wheat", "Wheat"),
        ("Maize", "Maize"), ("Corn", "Maize"), ("Cotton", "Cotton"),
        ("Sugarcane", "Sugarcane"), ("Potato", "Potato"), ("Onion", "Onion"),
        ("Chilli", "Chilli"), ("Chili", "Chilli"), ("Tomato", "Tomato"),
        ("Chickpea", "Chickpea"), ("Bengal Gram", "Bengal Gram"), ("Gram", "Chickpea"),
        ("Mustard", "Mustard"), ("Soybean", "Soybean"), ("Soya", "Soybean"),
        ("Groundnut", "Groundnut"), ("Peanut", "Groundnut"), ("Jowar", "Jowar"),
        ("Sorghum", "Jowar"), ("Bajra", "Bajra"), ("Garlic", "Garlic"),
        ("Turmeric", "Turmeric"), ("Banana", "Banana"), ("Grapes", "Grapes"),
        ("Mango", "Mango"), ("Vegetables", "Vegetables"), ("Pulses", "Pulses"),
        ("Oilseeds", "Oilseeds")
    ]
    for needle, standard in known_crops:
        if re.search(rf"\b{re.escape(needle)}\b", clean, re.IGNORECASE):
            if standard not in crops:
                crops.append(standard)
    
    if is_all and "All Crops" not in crops:
        crops.insert(0, "All Crops")
        
    return crops or ["All Crops"]


def _build_search_terms(name: str, generic: str, active: str, company: str) -> List[str]:
    """Creates a list of normalized lowercase keywords for matching shop inventory."""
    terms = set()
    raw = f"{name} {generic} {active} {company}".lower()
    # Split by slashes, dashes, spaces, parens
    tokens = re.split(r"[\s\(\)\/\,\-\:\+]+", raw)
    for t in tokens:
        t_clean = re.sub(r"[^\w]", "", t).strip()
        if len(t_clean) >= 3 and not t_clean.isdigit():
            terms.add(t_clean)
            
    # Add exact compounds
    name_l = name.lower()
    if "urea" in name_l: terms.add("urea")
    if "dap" in name_l: terms.add("dap")
    if "mop" in name_l: terms.add("mop"); terms.add("potash")
    if "npk" in name_l: terms.add("npk")
    if "ssp" in name_l: terms.add("ssp")
    if "zinc" in name_l: terms.add("zinc")
    if "boron" in name_l or "borax" in name_l: terms.add("boron"); terms.add("borax")
    if "mancozeb" in name_l: terms.add("mancozeb")
    if "carbendazim" in name_l: terms.add("carbendazim")
    if "imidacloprid" in name_l: terms.add("imidacloprid")
    if "pendimethalin" in name_l: terms.add("pendimethalin")
    if "chlorpyrifos" in name_l: terms.add("chlorpyrifos")
    if "emamectin" in name_l: terms.add("emamectin")
    
    return list(terms)


def load_catalog(force_reload: bool = False) -> List[Dict[str, Any]]:
    """Loads all 7 Excel files and produces the 181 normalized product records."""
    global _CATALOG_CACHE, _CATALOG_BY_ID, _CATALOG_BY_CATEGORY
    if _CATALOG_CACHE is not None and not force_reload:
        return _CATALOG_CACHE

    products: List[Dict[str, Any]] = []
    if not DATA_DIR or not os.path.exists(DATA_DIR):
        logger.warning(f"Data directory not found for agri catalog: {DATA_DIR}")
        _CATALOG_CACHE = []
        return []

    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl is required to load the Excel agri catalog.")
        _CATALOG_CACHE = []
        return []

    # 1. Fertilizers (45 items)
    fert_file = os.path.join(DATA_DIR, "fertilizers.xlsx")
    if os.path.exists(fert_file):
        try:
            wb = openpyxl.load_workbook(fert_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Product Name / Brand"))
                gen_name = _clean_str(d.get("Generic / Chemical Name"))
                cat = _clean_str(d.get("Category"))
                comp = _clean_str(d.get("Nutrient Composition (N-P-K-S + Others)"))
                company = _clean_str(d.get("Manufacturer / Company"))
                timing = _clean_str(d.get("Application Timing"))
                crops = _extract_crops_list(_clean_str(d.get("Target Crops")))
                search_terms = _build_search_terms(p_name, gen_name, comp, company)
                
                products.append({
                    "id": f"FERT_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Fertilizer",
                    "sub_category": cat or "Fertilizer",
                    "company": company,
                    "composition": comp,
                    "fco_grade": _clean_str(d.get("FCO Grade / Specification")),
                    "form": _clean_str(d.get("Form")),
                    "packaging": _clean_str(d.get("Packaging")),
                    "unit": _clean_str(d.get("Packaging")) or "50kg Bag",
                    "mrp": _clean_str(d.get("MRP (Approx ₹)")),
                    "subsidy_status": _clean_str(d.get("Subsidy Status")),
                    "target_crops": crops,
                    "purpose": _clean_str(d.get("Purpose / Use")),
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage (per acre)")),
                    "application_timing": timing,
                    "application_method": _clean_str(d.get("Application Method")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "precautions": _clean_str(d.get("Precautions / Safety Notes")),
                    "search_terms": search_terms,
                    "color": "text-emerald-700 bg-emerald-100 dark:text-emerald-400 dark:bg-emerald-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading fertilizers.xlsx: {e}")

    # 2. Insecticides (25 items)
    ins_file = os.path.join(DATA_DIR, "insecticides.xlsx")
    if os.path.exists(ins_file):
        try:
            wb = openpyxl.load_workbook(ins_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Brand / Trade Name"))
                gen_name = _clean_str(d.get("Active Ingredient (Technical)"))
                company = _clean_str(d.get("Company"))
                crops = _extract_crops_list(_clean_str(d.get("Registered Crops")))
                timing = _clean_str(d.get("Application Timing"))
                search_terms = _build_search_terms(p_name, gen_name, "", company)
                
                products.append({
                    "id": f"INSECT_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Insecticide",
                    "sub_category": _clean_str(d.get("Chemical Group / MOA Class")),
                    "company": company,
                    "formulation": _clean_str(d.get("Formulation")),
                    "composition": gen_name,
                    "unit": "1L / 500ml",
                    "target_pests": _clean_str(d.get("Target Pests")),
                    "target_crops": crops,
                    "mode_of_action": _clean_str(d.get("Mode of Action")),
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage (per acre)")),
                    "spray_volume": _clean_str(d.get("Spray Volume (water/acre)")),
                    "application_timing": timing,
                    "phi_days": _clean_str(d.get("PHI (days)")),
                    "toxicity_class": _clean_str(d.get("Toxicity Class (WHO)")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "precautions": _clean_str(d.get("Precautions / Safety")),
                    "search_terms": search_terms,
                    "color": "text-rose-700 bg-rose-100 dark:text-rose-400 dark:bg-rose-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading insecticides.xlsx: {e}")

    # 3. Fungicides (18 items)
    fung_file = os.path.join(DATA_DIR, "fungicides.xlsx")
    if os.path.exists(fung_file):
        try:
            wb = openpyxl.load_workbook(fung_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Brand / Trade Name"))
                gen_name = _clean_str(d.get("Active Ingredient (Technical)"))
                company = _clean_str(d.get("Company"))
                crops = _extract_crops_list(_clean_str(d.get("Registered Crops")))
                timing = _clean_str(d.get("Application Timing"))
                search_terms = _build_search_terms(p_name, gen_name, "", company)
                
                products.append({
                    "id": f"FUNG_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Fungicide",
                    "sub_category": _clean_str(d.get("Chemical Group / MOA (FRAC)")),
                    "company": company,
                    "formulation": _clean_str(d.get("Formulation")),
                    "composition": gen_name,
                    "unit": "1kg / 500g",
                    "target_diseases": _clean_str(d.get("Target Diseases")),
                    "target_crops": crops,
                    "mode_of_action": _clean_str(d.get("Mode of Action")),
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage (per acre)")),
                    "spray_volume": _clean_str(d.get("Spray Volume")),
                    "application_timing": timing,
                    "phi_days": _clean_str(d.get("PHI (days)")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "precautions": _clean_str(d.get("Precautions / Safety")),
                    "search_terms": search_terms,
                    "color": "text-purple-700 bg-purple-100 dark:text-purple-400 dark:bg-purple-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading fungicides.xlsx: {e}")

    # 4. Herbicides (17 items)
    herb_file = os.path.join(DATA_DIR, "herbicides_weedicides.xlsx")
    if os.path.exists(herb_file):
        try:
            wb = openpyxl.load_workbook(herb_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Brand / Trade Name"))
                gen_name = _clean_str(d.get("Active Ingredient (Technical)"))
                company = _clean_str(d.get("Company"))
                crops = _extract_crops_list(_clean_str(d.get("Registered Crops")))
                timing = _clean_str(d.get("Application Timing"))
                search_terms = _build_search_terms(p_name, gen_name, "", company)
                
                products.append({
                    "id": f"HERB_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Herbicide",
                    "sub_category": _clean_str(d.get("Type (Pre/Post)")),
                    "company": company,
                    "formulation": _clean_str(d.get("Formulation")),
                    "composition": gen_name,
                    "unit": "1L / 500ml",
                    "target_weeds": _clean_str(d.get("Target Weeds")),
                    "target_crops": crops,
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage (per acre)")),
                    "application_timing": timing,
                    "selectivity": _clean_str(d.get("Selectivity")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "precautions": _clean_str(d.get("Precautions / Safety")),
                    "search_terms": search_terms,
                    "color": "text-amber-700 bg-amber-100 dark:text-amber-400 dark:bg-amber-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading herbicides_weedicides.xlsx: {e}")

    # 5. Seeds (46 items)
    seed_file = os.path.join(DATA_DIR, "seeds.xlsx")
    if os.path.exists(seed_file):
        try:
            wb = openpyxl.load_workbook(seed_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                crop_name = _clean_str(d.get("Crop"))
                variety = _clean_str(d.get("Variety / Hybrid Name"))
                company = _clean_str(d.get("Company / Developer"))
                full_name = f"{crop_name} Seed ({variety})"
                search_terms = _build_search_terms(full_name, variety, crop_name, company)
                
                products.append({
                    "id": f"SEED_{sno}",
                    "name": full_name,
                    "generic_name": variety,
                    "category": "Seeds",
                    "sub_category": _clean_str(d.get("Type (Hybrid/OPV/HYV/Certified)")),
                    "crop": crop_name,
                    "company": company,
                    "unit": _clean_str(d.get("Seed Rate (per acre)")) or "Bags",
                    "target_crops": [crop_name],
                    "traits": _clean_str(d.get("Key Traits")),
                    "sowing_season": _clean_str(d.get("Sowing Season")),
                    "maturity_days": _clean_str(d.get("Maturity (Days)")),
                    "yield_potential": _clean_str(d.get("Yield Potential (q/acre)")),
                    "suitable_regions": _clean_str(d.get("Suitable Regions / States")),
                    "mrp": _clean_str(d.get("Seed Price (Approx ₹)")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "search_terms": search_terms,
                    "color": "text-yellow-700 bg-yellow-100 dark:text-yellow-400 dark:bg-yellow-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading seeds.xlsx: {e}")

    # 6. Bio-Inputs & PGRs (20 items)
    bio_file = os.path.join(DATA_DIR, "bio_inputs_pgrs.xlsx")
    if os.path.exists(bio_file):
        try:
            wb = openpyxl.load_workbook(bio_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Product Name"))
                gen_name = _clean_str(d.get("Active Organism / Ingredient"))
                company = _clean_str(d.get("Company / Source"))
                crops = _extract_crops_list(_clean_str(d.get("Target Crops")))
                timing = _clean_str(d.get("Application Timing"))
                search_terms = _build_search_terms(p_name, gen_name, "", company)
                
                products.append({
                    "id": f"BIO_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Bio-Inputs & PGRs",
                    "sub_category": _clean_str(d.get("Category")),
                    "company": company,
                    "composition": gen_name,
                    "unit": "1kg / 1L",
                    "purpose": _clean_str(d.get("Target Purpose")),
                    "target_crops": crops,
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage")),
                    "application_method": _clean_str(d.get("Application Method")),
                    "application_timing": timing,
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "search_terms": search_terms,
                    "color": "text-teal-700 bg-teal-100 dark:text-teal-400 dark:bg-teal-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading bio_inputs_pgrs.xlsx: {e}")

    # 7. Adjuvants & Others (10 items)
    adj_file = os.path.join(DATA_DIR, "adjuvants_others.xlsx")
    if os.path.exists(adj_file):
        try:
            wb = openpyxl.load_workbook(adj_file, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            for r in rows[1:]:
                if not any(r): continue
                d = dict(zip(headers, [r[i] if i < len(r) else None for i in range(len(headers))]))
                sno = _clean_str(d.get("S.No"))
                p_name = _clean_str(d.get("Product Name"))
                gen_name = _clean_str(d.get("Active Ingredient / Composition"))
                company = _clean_str(d.get("Company"))
                crops = _extract_crops_list(_clean_str(d.get("Target Crops / Usage")))
                search_terms = _build_search_terms(p_name, gen_name, "", company)
                
                products.append({
                    "id": f"ADJ_{sno}",
                    "name": p_name,
                    "generic_name": gen_name,
                    "category": "Adjuvants & Others",
                    "sub_category": _clean_str(d.get("Category")),
                    "company": company,
                    "composition": gen_name,
                    "unit": "500ml / 1L",
                    "purpose": _clean_str(d.get("Purpose")),
                    "target_crops": crops,
                    "recommended_dosage": _clean_str(d.get("Recommended Dosage")),
                    "application_method": _clean_str(d.get("Application Method")),
                    "key_benefits": _clean_str(d.get("Key Benefits")),
                    "search_terms": search_terms,
                    "color": "text-indigo-700 bg-indigo-100 dark:text-indigo-400 dark:bg-indigo-950/60"
                })
        except Exception as e:
            logger.error(f"Error loading adjuvants_others.xlsx: {e}")

    _CATALOG_CACHE = products
    _CATALOG_BY_ID = {p["id"]: p for p in products}
    
    _CATALOG_BY_CATEGORY = {}
    for p in products:
        c = p["category"]
        if c not in _CATALOG_BY_CATEGORY:
            _CATALOG_BY_CATEGORY[c] = []
        _CATALOG_BY_CATEGORY[c].append(p)

    logger.info(f"Loaded {len(products)} total products from Excel agri catalog across 7 files.")
    return products


def get_catalog() -> List[Dict[str, Any]]:
    """Returns the cached product catalog."""
    if _CATALOG_CACHE is None:
        return load_catalog()
    return _CATALOG_CACHE


# ---------------------------------------------------------------------------
# 2. CROP CANONICAL NORMALIZATION & LIFE CYCLE DEFINITIONS
# ---------------------------------------------------------------------------

CROP_CANONICAL_MAP = {
    "paddy": "Paddy (Rice)", "rice": "Paddy (Rice)", "wheat": "Wheat",
    "cotton": "Cotton", "maize": "Maize", "corn": "Maize", "sweet corn": "Maize",
    "chilli": "Chilli", "chili": "Chilli", "sugarcane": "Sugarcane",
    "chickpea": "Chickpea", "bengal gram": "Bengal Gram", "gram": "Chickpea",
    "potato": "Potato", "mustard": "Mustard", "onion": "Onion",
    "groundnut": "Groundnut", "peanut": "Groundnut", "soybean": "Soybean",
    "soya": "Soybean", "jowar": "Jowar", "sorghum": "Jowar", "tomato": "Tomato",
    "bajra": "Bajra", "pearl millet": "Bajra"
}

def normalize_crop_name(raw: str) -> str:
    n = (raw or "").strip().lower()
    if "bengal gram" in n: return "Bengal Gram"
    for key, canonical in CROP_CANONICAL_MAP.items():
        if key in n: return canonical
    return (raw or "Other").strip().title()


# Detailed Agronomic Lifecycle Stages for Major Indian Crops
CROP_LIFECYCLES: Dict[str, Dict[str, Any]] = {
    "Paddy (Rice)": {
        "duration_days": 130,
        "stages": [
            {
                "name": "Basal Sowing & Nursery",
                "start_day": 0, "end_day": 20,
                "description": "Nursery preparation, basal fertilizer placement before transplanting",
                "recommended_inputs": ["DAP", "Zinc Sulphate", "Pendimethalin", "Trichoderma viride"],
                "urgency": "high"
            },
            {
                "name": "Transplanting & Early Vegetative",
                "start_day": 21, "end_day": 40,
                "description": "Root establishment, initial split top-dress of Nitrogen",
                "recommended_inputs": ["Urea", "Zinc Sulphate", "Pretilachlor"],
                "urgency": "high"
            },
            {
                "name": "Active Tillering",
                "start_day": 41, "end_day": 65,
                "description": "Maximum tiller count determination; second Nitrogen + Potash top-dress",
                "recommended_inputs": ["Urea", "MOP Potash", "Chlorantraniliprole", "Mancozeb"],
                "urgency": "critical"
            },
            {
                "name": "Panicle Initiation & Booting",
                "start_day": 66, "end_day": 90,
                "description": "Head emergence and flowering; foliar nutrient boost & pest protection",
                "recommended_inputs": ["Nano Urea", "NPK 19:19:19", "Tricyclazole", "Imidacloprid"],
                "urgency": "high"
            },
            {
                "name": "Grain Filling & Milking",
                "start_day": 91, "end_day": 115,
                "description": "Starch deposition into grains; Potash foliar spray to prevent chaffiness",
                "recommended_inputs": ["Potassium Sulphate (SOP)", "Azoxystrobin"],
                "urgency": "normal"
            },
            {
                "name": "Maturity & Pre-Harvest",
                "start_day": 116, "end_day": 130,
                "description": "Golden grain ripening; cease chemical applications and drain water",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Wheat": {
        "duration_days": 125,
        "stages": [
            {
                "name": "Basal Sowing & Germination",
                "start_day": 0, "end_day": 20,
                "description": "Seedbed preparation with full phosphorus and potash basal placement",
                "recommended_inputs": ["DAP", "MOP Potash", "Zinc Sulphate", "Pendimethalin"],
                "urgency": "high"
            },
            {
                "name": "Crown Root Initiation (CRI)",
                "start_day": 21, "end_day": 35,
                "description": "First irrigation & critical first top-dress Urea application",
                "recommended_inputs": ["Urea", "2,4-D", "Metsulfuron"],
                "urgency": "critical"
            },
            {
                "name": "Tillering & Jointing",
                "start_day": 36, "end_day": 65,
                "description": "Stem elongation and active tiller formation; second Urea top-dress",
                "recommended_inputs": ["Urea", "NPK 19:19:19", "Sulfosulfuron"],
                "urgency": "high"
            },
            {
                "name": "Booting & Heading",
                "start_day": 66, "end_day": 90,
                "description": "Ear head emergence; protective rust fungicide and foliar nutrition",
                "recommended_inputs": ["Tebuconazole", "Propiconazole", "Nano Urea"],
                "urgency": "high"
            },
            {
                "name": "Grain Milking & Dough",
                "start_day": 91, "end_day": 110,
                "description": "Grain development and starch synthesis; guard against terminal heat",
                "recommended_inputs": ["Potassium Sulphate (SOP)"],
                "urgency": "normal"
            },
            {
                "name": "Ripening & Harvest",
                "start_day": 111, "end_day": 125,
                "description": "Golden straw maturity; prepare harvesting machinery",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Cotton": {
        "duration_days": 170,
        "stages": [
            {
                "name": "Sowing & Seedling",
                "start_day": 0, "end_day": 25,
                "description": "Basal NPK application & pre-emergence weed control",
                "recommended_inputs": ["NPK 10:26:26", "DAP", "Pendimethalin", "Imidacloprid"],
                "urgency": "high"
            },
            {
                "name": "Square Formation (Vegetative)",
                "start_day": 26, "end_day": 60,
                "description": "Branching and square (flower bud) initiation; top-dress Nitrogen & Zinc",
                "recommended_inputs": ["Urea", "Zinc Sulphate", "Acetamiprid", "Borax"],
                "urgency": "critical"
            },
            {
                "name": "Flowering & Boll Setting",
                "start_day": 61, "end_day": 110,
                "description": "Peak boll load formation; high Potassium and Nitrogen demand; bollworm protection",
                "recommended_inputs": ["Urea", "MOP Potash", "Emamectin Benzoate", "Chlorantraniliprole"],
                "urgency": "critical"
            },
            {
                "name": "Boll Development & Bursting",
                "start_day": 111, "end_day": 145,
                "description": "Lint maturation; foliar nutrition to prevent square/boll drop",
                "recommended_inputs": ["NPK 19:19:19", "Magnesium Sulphate"],
                "urgency": "high"
            },
            {
                "name": "Harvest & Multi-Picking",
                "start_day": 146, "end_day": 170,
                "description": "First and subsequent pickings of mature bolls",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Maize": {
        "duration_days": 110,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 15,
                "description": "Basal phosphorus, zinc and pre-emergence weed control",
                "recommended_inputs": ["DAP", "Zinc Sulphate", "Atrazine"],
                "urgency": "high"
            },
            {
                "name": "Knee-High (Vegetative)",
                "start_day": 16, "end_day": 35,
                "description": "Rapid vegetative growth; first Urea top-dress and Fall Armyworm monitoring",
                "recommended_inputs": ["Urea", "Chlorantraniliprole", "Spinetoram"],
                "urgency": "critical"
            },
            {
                "name": "Tasseling & Silking",
                "start_day": 36, "end_day": 65,
                "description": "Flowering & cob initiation; second top-dress Nitrogen + Potash",
                "recommended_inputs": ["Urea", "MOP Potash", "Borax"],
                "urgency": "critical"
            },
            {
                "name": "Cob Filling & Milking",
                "start_day": 66, "end_day": 90,
                "description": "Grain setting and kernel filling; foliar nutrient support",
                "recommended_inputs": ["NPK 19:19:19", "Nano Urea"],
                "urgency": "normal"
            },
            {
                "name": "Maturity & Harvest",
                "start_day": 91, "end_day": 110,
                "description": "Husk drying and black-layer maturity",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Potato": {
        "duration_days": 105,
        "stages": [
            {
                "name": "Planting & Basal",
                "start_day": 0, "end_day": 20,
                "description": "High basal phosphate and potash for tuber foundation; seed tuber treatment",
                "recommended_inputs": ["DAP", "MOP Potash", "NPK 10:26:26", "Mancozeb"],
                "urgency": "critical"
            },
            {
                "name": "Emergence & Earthing Up",
                "start_day": 21, "end_day": 40,
                "description": "Canopy establishment; earthing up with Urea top-dress and weed control",
                "recommended_inputs": ["Urea", "Metribuzin", "Zinc Sulphate"],
                "urgency": "critical"
            },
            {
                "name": "Tuber Initiation & Bulking",
                "start_day": 41, "end_day": 70,
                "description": "Stolon tip swelling and rapid bulking; high potash & late blight prevention",
                "recommended_inputs": ["MOP Potash", "Mancozeb", "Cymoxanil", "Borax"],
                "urgency": "critical"
            },
            {
                "name": "Tuber Maturation & Sizing",
                "start_day": 71, "end_day": 90,
                "description": "Skin hardening and sizing; foliar Potash application",
                "recommended_inputs": ["Potassium Sulphate (SOP)", "Dimethomorph"],
                "urgency": "high"
            },
            {
                "name": "Dehaulming & Digging",
                "start_day": 91, "end_day": 105,
                "description": "Cut haulms 10 days before digging for skin curing",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Onion": {
        "duration_days": 130,
        "stages": [
            {
                "name": "Transplanting & Basal",
                "start_day": 0, "end_day": 25,
                "description": "Seedling transplanting with basal DAP, Potash, and bentonite Sulphur",
                "recommended_inputs": ["DAP", "MOP Potash", "Single Super Phosphate (SSP)", "Pendimethalin"],
                "urgency": "high"
            },
            {
                "name": "Vegetative & Foliage Growth",
                "start_day": 26, "end_day": 55,
                "description": "Leaf production (more leaves = bigger bulb); top-dress Nitrogen & Thrips control",
                "recommended_inputs": ["Urea", "Ammonium Sulphate", "Fipronil", "Silicone Spreader"],
                "urgency": "critical"
            },
            {
                "name": "Bulb Initiation & Enlargement",
                "start_day": 56, "end_day": 90,
                "description": "Bulb swelling; high Potash and Sulphur demand; Purple blotch prevention",
                "recommended_inputs": ["MOP Potash", "Mancozeb", "Borax", "Silicone Spreader"],
                "urgency": "critical"
            },
            {
                "name": "Bulb Maturation & Neck Fall",
                "start_day": 91, "end_day": 115,
                "description": "Cease nitrogen; foliar 0-0-50 for firmness and storage quality",
                "recommended_inputs": ["Potassium Sulphate (SOP)"],
                "urgency": "normal"
            },
            {
                "name": "Harvest & Field Curing",
                "start_day": 116, "end_day": 130,
                "description": "Lift bulbs after 50% tops fall; field curing in shade",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Chilli": {
        "duration_days": 160,
        "stages": [
            {
                "name": "Transplanting & Establishment",
                "start_day": 0, "end_day": 25,
                "description": "Basal complex fertilizer NPK 20:20:0:13 and root dip protection",
                "recommended_inputs": ["NPK 20:20:0:13", "DAP", "Trichoderma viride"],
                "urgency": "high"
            },
            {
                "name": "Vegetative Branching",
                "start_day": 26, "end_day": 55,
                "description": "Canopy development; top-dress Nitrogen and thrips/mite protection",
                "recommended_inputs": ["Urea", "Zinc Sulphate", "Diafenthiuron", "Imidacloprid"],
                "urgency": "critical"
            },
            {
                "name": "Flowering & Fruit Setting",
                "start_day": 56, "end_day": 90,
                "description": "Heavy blooming and fruit set; prevent flower drop; Anthracnose fungicide",
                "recommended_inputs": ["NPK 19:19:19", "Borax", "Azoxystrobin", "Mancozeb"],
                "urgency": "critical"
            },
            {
                "name": "Fruit Development & Pickings",
                "start_day": 91, "end_day": 140,
                "description": "Continuous fruit picking; split potassium and calcium feeding",
                "recommended_inputs": ["Calcium Ammonium Nitrate (CAN)", "MOP Potash"],
                "urgency": "high"
            },
            {
                "name": "Final Red Ripening & Harvest",
                "start_day": 141, "end_day": 160,
                "description": "Harvesting red chillies for drying",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Chickpea": {
        "duration_days": 105,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "Basal phosphorus for root nodulation and rhizobium colonization",
                "recommended_inputs": ["DAP", "Single Super Phosphate (SSP)", "Trichoderma viride"],
                "urgency": "high"
            },
            {
                "name": "Vegetative & Branching",
                "start_day": 21, "end_day": 45,
                "description": "Nipping top shoots to stimulate lateral fruiting branches; light weeding",
                "recommended_inputs": ["Pendimethalin", "NPK 19:19:19"],
                "urgency": "normal"
            },
            {
                "name": "Flowering & Pod Formation",
                "start_day": 46, "end_day": 75,
                "description": "CRITICAL stage for Pod Borer (Helicoverpa armigera) control & foliar DAP",
                "recommended_inputs": ["Emamectin Benzoate", "Chlorantraniliprole", "Nano DAP"],
                "urgency": "critical"
            },
            {
                "name": "Pod Filling & Maturation",
                "start_day": 76, "end_day": 95,
                "description": "Grain filling; stop irrigation to prevent root rot",
                "recommended_inputs": ["NPK 19:19:19"],
                "urgency": "normal"
            },
            {
                "name": "Harvest",
                "start_day": 96, "end_day": 105,
                "description": "Brown straw drying and threshing",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Bengal Gram": {
        "duration_days": 105,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "Phosphorus basal placement for nitrogen-fixing nodules",
                "recommended_inputs": ["DAP", "Single Super Phosphate (SSP)", "Trichoderma viride"],
                "urgency": "high"
            },
            {
                "name": "Vegetative & Branching",
                "start_day": 21, "end_day": 45,
                "description": "Branching and foliage development",
                "recommended_inputs": ["NPK 19:19:19"],
                "urgency": "normal"
            },
            {
                "name": "Flowering & Pod Borer Defense",
                "start_day": 46, "end_day": 75,
                "description": "Defend against Helicoverpa pod borer and spray foliar phosphorus",
                "recommended_inputs": ["Emamectin Benzoate", "Chlorantraniliprole", "Nano DAP"],
                "urgency": "critical"
            },
            {
                "name": "Pod Filling & Maturation",
                "start_day": 76, "end_day": 95,
                "description": "Grain sizing inside pods",
                "recommended_inputs": [],
                "urgency": "normal"
            },
            {
                "name": "Harvest",
                "start_day": 96, "end_day": 105,
                "description": "Pod rattling and harvesting",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Mustard": {
        "duration_days": 115,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "High Sulphur requirement (essential for oil synthesis); SSP + Urea basal",
                "recommended_inputs": ["Single Super Phosphate (SSP)", "Urea", "Pendimethalin"],
                "urgency": "high"
            },
            {
                "name": "Rosette & Branching",
                "start_day": 21, "end_day": 45,
                "description": "First irrigation and split top-dress Urea",
                "recommended_inputs": ["Urea", "Zinc Sulphate"],
                "urgency": "critical"
            },
            {
                "name": "Flowering & Pod (Siliqua) Formation",
                "start_day": 46, "end_day": 75,
                "description": "Crucial mustard aphid protection and foliar Sulphur/Boron",
                "recommended_inputs": ["Dimethoate", "Imidacloprid", "Borax", "Mancozeb"],
                "urgency": "critical"
            },
            {
                "name": "Pod Filling & Oil Synthesis",
                "start_day": 76, "end_day": 100,
                "description": "Seed development and oil percentage accumulation",
                "recommended_inputs": ["NPK 19:19:19"],
                "urgency": "normal"
            },
            {
                "name": "Maturity & Harvest",
                "start_day": 101, "end_day": 115,
                "description": "Pods turning golden-yellow; threshing before pod shattering",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Soybean": {
        "duration_days": 100,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "DAP + SSP for phosphorus and sulphur; seed treatment with Rhizobium",
                "recommended_inputs": ["DAP", "Single Super Phosphate (SSP)", "Pendimethalin", "Trichoderma viride"],
                "urgency": "high"
            },
            {
                "name": "Vegetative Growth (V3-V5)",
                "start_day": 21, "end_day": 40,
                "description": "Post-emergence weed management & girdle beetle protection",
                "recommended_inputs": ["Imazethapyr", "Chlorantraniliprole"],
                "urgency": "high"
            },
            {
                "name": "Flowering & Pod Initiation (R1-R3)",
                "start_day": 41, "end_day": 65,
                "description": "Foliar nutrition & defoliator caterpillar control",
                "recommended_inputs": ["NPK 19:19:19", "Borax", "Emamectin Benzoate"],
                "urgency": "critical"
            },
            {
                "name": "Pod Filling & Grain Development",
                "start_day": 66, "end_day": 85,
                "description": "Rust and anthracnose preventive sprays",
                "recommended_inputs": ["Tebuconazole", "Mancozeb"],
                "urgency": "normal"
            },
            {
                "name": "Maturity & Harvest",
                "start_day": 86, "end_day": 100,
                "description": "Leaves yellow and drop; pods dry for combining",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Groundnut": {
        "duration_days": 120,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "Gypsum and SSP placement for calcium & sulphur required for pod walls",
                "recommended_inputs": ["Gypsum", "Single Super Phosphate (SSP)", "DAP"],
                "urgency": "high"
            },
            {
                "name": "Vegetative & Flowering",
                "start_day": 21, "end_day": 40,
                "description": "Yellow flowers emerge; light earthing up and leaf miner control",
                "recommended_inputs": ["Urea", "Imidacloprid"],
                "urgency": "high"
            },
            {
                "name": "Pegging (Most Critical Stage)",
                "start_day": 41, "end_day": 70,
                "description": "Pegs penetrate soil to form pods; MUST apply second dose of Gypsum (100 kg/ac)",
                "recommended_inputs": ["Gypsum", "Mancozeb", "Carbendazim"],
                "urgency": "critical"
            },
            {
                "name": "Pod Development & Filling",
                "start_day": 71, "end_day": 95,
                "description": "Soil moisture maintenance; Tikka disease defense",
                "recommended_inputs": ["Tebuconazole", "NPK 19:19:19"],
                "urgency": "high"
            },
            {
                "name": "Maturity & Pod Lifting",
                "start_day": 96, "end_day": 120,
                "description": "Internal pod shell turning blackish-brown; harvesting & drying",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Sugarcane": {
        "duration_days": 360,
        "stages": [
            {
                "name": "Germination & Settling",
                "start_day": 0, "end_day": 45,
                "description": "Sett placement with basal NPK 10:26:26 and termite / early shoot borer control",
                "recommended_inputs": ["NPK 10:26:26", "DAP", "Chlorpyrifos"],
                "urgency": "high"
            },
            {
                "name": "Tillering & Formative",
                "start_day": 46, "end_day": 120,
                "description": "Heavy Nitrogen top-dress splits & earthing up to prevent shoot borer damage",
                "recommended_inputs": ["Urea", "Chlorantraniliprole", "Zinc Sulphate"],
                "urgency": "critical"
            },
            {
                "name": "Grand Growth & Cane Elongation",
                "start_day": 121, "end_day": 270,
                "description": "Rapid cane height and internode formation; Nitrogen + Potash top-dress",
                "recommended_inputs": ["Urea", "MOP Potash"],
                "urgency": "critical"
            },
            {
                "name": "Ripening & Sucrose Accumulation",
                "start_day": 271, "end_day": 360,
                "description": "Sugar synthesis in cane tissue; stop irrigation 2-3 weeks before cutting",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Tomato": {
        "duration_days": 130,
        "stages": [
            {
                "name": "Transplanting & Basal",
                "start_day": 0, "end_day": 20,
                "description": "Basal NPK 19:19:19 and soil amendment with neem cake & Trichoderma",
                "recommended_inputs": ["NPK 19:19:19", "DAP", "Trichoderma viride", "Neem Cake"],
                "urgency": "high"
            },
            {
                "name": "Vegetative Growth",
                "start_day": 21, "end_day": 45,
                "description": "Staking, suckering, first Nitrogen top-dress and Whitefly vector control",
                "recommended_inputs": ["Urea", "Imidacloprid", "Zinc Sulphate"],
                "urgency": "critical"
            },
            {
                "name": "Flowering & Fruit Setting",
                "start_day": 46, "end_day": 75,
                "description": "Calcium & Boron application to prevent Blossom End Rot; Blight prevention",
                "recommended_inputs": ["Calcium Ammonium Nitrate (CAN)", "Borax", "Mancozeb", "Azoxystrobin"],
                "urgency": "critical"
            },
            {
                "name": "Fruit Sizing & Continuous Picking",
                "start_day": 76, "end_day": 115,
                "description": "Potash feeding for firmness and color; regular fruit harvests",
                "recommended_inputs": ["Potassium Sulphate (SOP)", "NPK 19:19:19"],
                "urgency": "high"
            },
            {
                "name": "Final Pickings",
                "start_day": 116, "end_day": 130,
                "description": "Final crop clearance",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    },
    "Jowar": {
        "duration_days": 110,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 20,
                "description": "Basal DAP and Shoot Fly seed treatment",
                "recommended_inputs": ["DAP", "Zinc Sulphate", "Imidacloprid"],
                "urgency": "high"
            },
            {
                "name": "Vegetative (Knee-High)",
                "start_day": 21, "end_day": 45,
                "description": "Top-dress Urea and Stem Borer control",
                "recommended_inputs": ["Urea", "Chlorantraniliprole"],
                "urgency": "critical"
            },
            {
                "name": "Booting & Flowering",
                "start_day": 46, "end_day": 75,
                "description": "Panicle emergence and flowering",
                "recommended_inputs": ["NPK 19:19:19"],
                "urgency": "normal"
            },
            {
                "name": "Grain Filling & Dough",
                "start_day": 76, "end_day": 95,
                "description": "Starch accumulation in grains",
                "recommended_inputs": [],
                "urgency": "normal"
            },
            {
                "name": "Maturity & Harvest",
                "start_day": 96, "end_day": 110,
                "description": "Earheads turn brownish; cutting and threshing",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    }
}


def get_default_lifecycle(crop_name: str) -> Dict[str, Any]:
    """Fallback 4-stage lifecycle for crops not explicitly mapped."""
    return {
        "duration_days": 120,
        "stages": [
            {
                "name": "Basal Sowing",
                "start_day": 0, "end_day": 25,
                "description": "Field preparation, basal fertilizer placement, and germination",
                "recommended_inputs": ["DAP", "Urea", "Pendimethalin"],
                "urgency": "high"
            },
            {
                "name": "Vegetative Growth",
                "start_day": 26, "end_day": 60,
                "description": "Canopy establishment, branching, and top-dress nitrogen feeding",
                "recommended_inputs": ["Urea", "Zinc Sulphate", "NPK 19:19:19"],
                "urgency": "critical"
            },
            {
                "name": "Flowering & Fruiting",
                "start_day": 61, "end_day": 95,
                "description": "Reproductive phase, flower setting, pest and disease protection",
                "recommended_inputs": ["MOP Potash", "NPK 19:19:19", "Mancozeb"],
                "urgency": "high"
            },
            {
                "name": "Maturity & Harvest",
                "start_day": 96, "end_day": 120,
                "description": "Final ripening and harvest",
                "recommended_inputs": [],
                "urgency": "normal"
            }
        ]
    }


def get_crop_lifecycle(crop_name: str) -> Dict[str, Any]:
    canonical = normalize_crop_name(crop_name)
    return CROP_LIFECYCLES.get(canonical, get_default_lifecycle(canonical))


def calculate_crop_stage(
    sowing_date: Optional[datetime],
    crop_name: str,
    ref_date: Optional[datetime] = None
) -> Dict[str, Any]:
    """
    Given a sowing date and crop name, calculates:
    - Current growth stage
    - Days since sowing
    - Percentage through lifecycle
    - Days remaining to harvest
    - Products/fertilizers needed right now
    - Products needed in the next stage
    """
    if ref_date is None:
        ref_date = datetime.utcnow()

    canonical = normalize_crop_name(crop_name)
    lifecycle = get_crop_lifecycle(canonical)
    total_days = lifecycle["duration_days"]

    if sowing_date is None:
        # Default mid-cycle fallback
        sowing_date = ref_date - timedelta(days=total_days // 3)

    # Normalize timezone
    if hasattr(sowing_date, "tzinfo") and sowing_date.tzinfo is not None:
        sowing_date = sowing_date.replace(tzinfo=None)
    if hasattr(ref_date, "tzinfo") and ref_date.tzinfo is not None:
        ref_date = ref_date.replace(tzinfo=None)

    days_since_sowing = max(0, (ref_date - sowing_date).days)
    
    # Check if beyond duration
    is_overdue = days_since_sowing > total_days
    capped_days = min(days_since_sowing, total_days)
    progress_pct = round(min(100.0, (capped_days / total_days) * 100), 1)
    days_to_harvest = max(0, total_days - days_since_sowing)

    stages = lifecycle["stages"]
    current_stage_idx = 0
    for idx, stage in enumerate(stages):
        if stage["start_day"] <= days_since_sowing <= stage["end_day"]:
            current_stage_idx = idx
            break
        elif days_since_sowing > stage["end_day"]:
            current_stage_idx = idx

    current_stage = stages[current_stage_idx]
    next_stage = stages[current_stage_idx + 1] if current_stage_idx + 1 < len(stages) else None
    
    days_in_stage = days_since_sowing - current_stage["start_day"]
    days_to_next_stage = max(0, current_stage["end_day"] - days_since_sowing + 1) if next_stage else 0

    # Build stage timeline markers
    stage_timeline = []
    for idx, st in enumerate(stages):
        if idx < current_stage_idx:
            status = "completed"
        elif idx == current_stage_idx:
            status = "current"
        else:
            status = "upcoming"
        stage_timeline.append({
            "name": st["name"],
            "start_day": st["start_day"],
            "end_day": st["end_day"],
            "description": st["description"],
            "recommended_inputs": st["recommended_inputs"],
            "urgency": st["urgency"],
            "status": status
        })

    return {
        "crop_name": canonical,
        "days_since_sowing": days_since_sowing,
        "total_duration_days": total_days,
        "progress_pct": progress_pct,
        "days_to_harvest": days_to_harvest,
        "is_harvest_ready": days_to_harvest <= 14 or is_overdue,
        "stage_index": current_stage_idx,
        "total_stages": len(stages),
        "current_stage_name": current_stage["name"],
        "current_stage_description": current_stage["description"],
        "current_stage_urgency": current_stage["urgency"],
        "days_in_current_stage": days_in_stage,
        "days_to_next_stage": days_to_next_stage,
        "inputs_needed_now": current_stage["recommended_inputs"],
        "inputs_needed_next": next_stage["recommended_inputs"] if next_stage else [],
        "stage_timeline": stage_timeline
    }


# ---------------------------------------------------------------------------
# 3. CATALOG MATCHING ENGINE FOR REAL INVENTORY & STAGE RECOMMENDATIONS
# ---------------------------------------------------------------------------

def match_catalog_products_for_inputs(input_names: List[str]) -> List[Dict[str, Any]]:
    """Finds matching rich products from the 181-product Excel catalog."""
    catalog = get_catalog()
    matched = []
    seen_ids = set()
    
    for inp in input_names:
        inp_lower = inp.lower()
        for prod in catalog:
            if prod["id"] in seen_ids:
                continue
            # Match by search terms or name
            match = False
            for st in prod["search_terms"]:
                if st in inp_lower or inp_lower in st:
                    match = True
                    break
            if not match and (inp_lower in prod["name"].lower() or inp_lower in prod.get("generic_name", "").lower()):
                match = True
            
            if match:
                seen_ids.add(prod["id"])
                matched.append(prod)
                if len(matched) >= 8:  # avoid overwhelming list
                    break
        if len(matched) >= 8:
            break
            
    return matched
