"""
AgriFlow Database Loader V2 — Community Interlinked Edition
=============================================================
Reads agriflow_seed_data_v2.xlsx and inserts all users + profiles +
community-interlinked transactional data into the PostgreSQL database.

Run:  python seed_data/load_to_db_v2.py
"""

import os
import sys
import random
import asyncio
from datetime import datetime, timedelta, date
from pathlib import Path

# Add backend to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from dotenv import load_dotenv
load_dotenv(str(Path(__file__).resolve().parent.parent / ".env"))

from openpyxl import load_workbook
import bcrypt
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlmodel import SQLModel, select

# Import all models
from app.models.user import User, UserRole
from app.models.farmer import FarmerProfile, LandRecord
from app.models.crop import Crop, CropExpense, CropHarvest, CropSale
from app.models.shop import ShopProfile
from app.models.manufacturer import MillProfile, ManufacturerPurchase, ProductionBatch, ManufacturerSale
from app.models.customer import CustomerProfile, CustomerOrder, CustomerOrderItem
from app.models.trade import Product, ShopOrder, ShopOrderItem
from app.models.expense import ShopExpense
from app.models.shop_accounting import ShopAccountingExpense
from app.models.manufacturer_expense import ManufacturerExpense

random.seed(2026)

# ─── Crop data for generating realistic farming history ──────────────────────

CROPS_DATA = {
    "Kharif": [
        {"name": "Paddy", "type": "Cereal", "varieties": ["Sona Masuri", "BPT 5204", "IR-64", "Swarna", "Tellahamsa"], "yield_range": (15, 35), "price_range": (1800, 2500), "cost_per_acre": (12000, 22000)},
        {"name": "Maize", "type": "Cereal", "varieties": ["DHM-117", "Kaveri 50", "NK-6240"], "yield_range": (20, 40), "price_range": (1400, 2100), "cost_per_acre": (10000, 18000)},
        {"name": "Cotton", "type": "Commercial", "varieties": ["Bt Cotton", "Bunny", "Mallika"], "yield_range": (8, 18), "price_range": (5500, 7200), "cost_per_acre": (15000, 30000)},
        {"name": "Soybean", "type": "Oilseed", "varieties": ["JS-335", "JS-9560", "NRC-37"], "yield_range": (8, 15), "price_range": (3800, 5200), "cost_per_acre": (8000, 15000)},
        {"name": "Groundnut", "type": "Oilseed", "varieties": ["TMV-2", "JL-501", "Kadiri-6"], "yield_range": (10, 20), "price_range": (4500, 6500), "cost_per_acre": (12000, 20000)},
        {"name": "Chilli", "type": "Spice", "varieties": ["Teja", "Byadgi", "S-4"], "yield_range": (5, 15), "price_range": (8000, 18000), "cost_per_acre": (25000, 45000)},
        {"name": "Sugarcane", "type": "Commercial", "varieties": ["Co-86032", "CoC-671", "Co-0238"], "yield_range": (300, 500), "price_range": (280, 350), "cost_per_acre": (35000, 60000)},
        {"name": "Jowar", "type": "Cereal", "varieties": ["CSV-15", "CSH-16", "Maldandi"], "yield_range": (8, 18), "price_range": (2200, 3200), "cost_per_acre": (6000, 12000)},
        {"name": "Turmeric", "type": "Spice", "varieties": ["Erode Local", "Salem", "Rajapore"], "yield_range": (20, 40), "price_range": (6000, 12000), "cost_per_acre": (30000, 55000)},
        {"name": "Bajra", "type": "Cereal", "varieties": ["HHB-67", "ICTP-8203", "Raj-171"], "yield_range": (8, 16), "price_range": (1900, 2800), "cost_per_acre": (5000, 10000)},
    ],
    "Rabi": [
        {"name": "Wheat", "type": "Cereal", "varieties": ["PBW-343", "HD-2967", "Lok-1"], "yield_range": (15, 30), "price_range": (1900, 2600), "cost_per_acre": (10000, 18000)},
        {"name": "Chickpea", "type": "Pulse", "varieties": ["JG-11", "JAKI-9218", "Vijay"], "yield_range": (6, 12), "price_range": (4200, 5800), "cost_per_acre": (8000, 14000)},
        {"name": "Mustard", "type": "Oilseed", "varieties": ["Pusa Bold", "RH-749", "Bio-902"], "yield_range": (6, 12), "price_range": (4500, 6000), "cost_per_acre": (6000, 11000)},
        {"name": "Onion", "type": "Vegetable", "varieties": ["Nasik Red", "Bellary Red", "Pusa Ratnar"], "yield_range": (80, 150), "price_range": (800, 2500), "cost_per_acre": (30000, 50000)},
        {"name": "Potato", "type": "Vegetable", "varieties": ["Kufri Jyoti", "Kufri Pukhraj"], "yield_range": (80, 150), "price_range": (600, 1500), "cost_per_acre": (35000, 55000)},
        {"name": "Bengal Gram", "type": "Pulse", "varieties": ["JG-11", "JG-14", "KAK-2"], "yield_range": (5, 10), "price_range": (4500, 6200), "cost_per_acre": (7000, 12000)},
        {"name": "Lentil", "type": "Pulse", "varieties": ["IPL-316", "K-75", "Pant L-5"], "yield_range": (4, 8), "price_range": (3800, 5500), "cost_per_acre": (5000, 9000)},
    ],
}

EXPENSE_TYPES = {
    "Input": [
        ("Seed", "kg", 5, 50, 30, 200),
        ("Fertilizer (DAP)", "bag", 1, 5, 1100, 1400),
        ("Fertilizer (Urea)", "bag", 1, 4, 250, 320),
        ("Fertilizer (MOP)", "bag", 1, 3, 800, 1050),
        ("Pesticide", "liter", 0.5, 3, 300, 900),
        ("Herbicide", "liter", 0.5, 2, 250, 700),
        ("Micronutrients", "kg", 1, 5, 40, 80),
    ],
    "Labor": [
        ("Ploughing Labour", "days", 2, 5, 400, 700),
        ("Sowing Labour", "days", 1, 3, 350, 600),
        ("Weeding Labour", "days", 2, 6, 300, 550),
        ("Spraying Labour", "days", 1, 3, 350, 500),
        ("Harvesting Labour", "days", 3, 8, 400, 700),
    ],
    "Machinery": [
        ("Tractor Hire", "hours", 3, 10, 500, 1000),
        ("Rotavator", "hours", 2, 5, 600, 1200),
        ("Harvester Hire", "hours", 2, 6, 800, 1500),
    ],
    "Irrigation": [
        ("Borewell Electricity", "months", 1, 5, 800, 2500),
        ("Diesel Pump Fuel", "liters", 10, 50, 85, 95),
    ],
}

EXPENSE_STAGES = ["Sowing", "Germination", "Vegetative", "Flowering", "Fruiting", "Harvesting", "Post-Harvest"]

SHOP_PRODUCTS_DATA = [
    {"name": "DAP Fertilizer", "category": "fertilizer", "brand": "IFFCO", "unit": "bag", "qty_per_unit": 50, "cost": 1150, "price": 1350},
    {"name": "Urea", "category": "fertilizer", "brand": "IFFCO", "unit": "bag", "qty_per_unit": 50, "cost": 266, "price": 300},
    {"name": "MOP Potash", "category": "fertilizer", "brand": "IPL", "unit": "bag", "qty_per_unit": 50, "cost": 850, "price": 1000},
    {"name": "NPK 20-20-0", "category": "fertilizer", "brand": "Coromandel", "unit": "bag", "qty_per_unit": 50, "cost": 1100, "price": 1280},
    {"name": "Imidacloprid 17.8 SL", "category": "pesticide", "brand": "Bayer", "unit": "liter", "qty_per_unit": 1, "cost": 800, "price": 1050},
    {"name": "Chlorpyriphos 20 EC", "category": "pesticide", "brand": "Dhanuka", "unit": "liter", "qty_per_unit": 1, "cost": 350, "price": 480},
    {"name": "Carbendazim 50 WP", "category": "pesticide", "brand": "BASF", "unit": "packet", "qty_per_unit": 0.5, "cost": 180, "price": 260},
    {"name": "Paddy Seed BPT-5204", "category": "seeds", "brand": "NSC", "unit": "kg", "qty_per_unit": 10, "cost": 45, "price": 65},
    {"name": "Maize Seed Hybrid", "category": "seeds", "brand": "Kaveri", "unit": "kg", "qty_per_unit": 5, "cost": 180, "price": 250},
    {"name": "Neem Oil", "category": "pesticide", "brand": "Multiplex", "unit": "liter", "qty_per_unit": 1, "cost": 200, "price": 300},
    {"name": "Vermicompost", "category": "fertilizer", "brand": "Organic Gold", "unit": "bag", "qty_per_unit": 50, "cost": 300, "price": 450},
    {"name": "Zinc Sulphate", "category": "fertilizer", "brand": "Tata", "unit": "kg", "qty_per_unit": 25, "cost": 40, "price": 55},
]

# ─── Password hashing ────────────────────────────────────────────────────────
_hash_cache = {}
def get_hash(password: str) -> str:
    if password not in _hash_cache:
        salt = bcrypt.gensalt()
        _hash_cache[password] = bcrypt.hashpw(password.encode('utf-8'), salt).decode('utf-8')
    return _hash_cache[password]


# ─── Database setup ──────────────────────────────────────────────────────────
DATABASE_URL = os.getenv("DATABASE_URL", "")
if not DATABASE_URL:
    print("ERROR: DATABASE_URL not set in .env")
    sys.exit(1)

engine = create_async_engine(DATABASE_URL, echo=False, pool_size=10, max_overflow=20)
async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)


def rand_date_between(start: datetime, end: datetime) -> datetime:
    delta = (end - start).days
    if delta <= 0:
        return start
    return start + timedelta(days=random.randint(0, delta))


# ═══════════════════════════════════════════════════════════════════════════════
#                          LOAD DATA
# ═══════════════════════════════════════════════════════════════════════════════

async def load_all():
    xlsx_path = os.path.join(os.path.dirname(__file__), "agriflow_seed_data_v2.xlsx")
    if not os.path.exists(xlsx_path):
        print(f"ERROR: {xlsx_path} not found. Run generate_excel_v2.py first.")
        sys.exit(1)

    print("Loading Excel workbook (V2)...")
    wb = load_workbook(xlsx_path, read_only=True)

    print("Pre-computing password hashes...")
    farmer_hash = get_hash("Farmer@123")
    shop_hash = get_hash("Shop@123")
    mill_hash = get_hash("Mill@123")
    customer_hash = get_hash("Customer@123")

    now = datetime.utcnow()
    batch_size = 500

    # ─── Mappings from Excel idx → DB user_id ────────────────────────────────
    farmer_idx_to_uid = {}   # farmer_idx → user_id
    shop_idx_to_uid = {}     # shop_idx → user_id
    mill_idx_to_uid = {}     # mill_idx → user_id
    customer_idx_to_uid = {} # cust_idx → user_id

    # Also track product IDs per shop/mill
    shop_product_ids = {}    # shop_user_id → [(product_id, product_data), ...]

    # ─── 1. Load Farmers ──────────────────────────────────────────────────────
    ws = wb["farmers"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_farmers = len(rows)
    print(f"\n{'='*60}")
    print(f"  LOADING {total_farmers:,} FARMERS")
    print(f"{'='*60}")

    for batch_start in range(0, total_farmers, batch_size):
        batch_end = min(batch_start + batch_size, total_farmers)
        batch = rows[batch_start:batch_end]

        async with async_session() as session:
            for row in batch:
                (farmer_idx, email, phone, password, full_name, farmer_id,
                 father_name, gender, relation_type,
                 house_no, street, village, mandal, district, state, pincode,
                 region_name, region_idx,
                 total_area, aadhaar_last_4, bank_name, account_number, ifsc_code,
                 primary_crop, secondary_crop,
                 latitude, longitude,
                 land_serial_1, land_area_1, land_serial_2, land_area_2) = row

                user = User(
                    email=email,
                    phone_number=str(phone),
                    full_name=full_name,
                    role=UserRole.FARMER,
                    is_active=True,
                    hashed_password=farmer_hash,
                )
                session.add(user)
                await session.flush()

                profile = FarmerProfile(
                    user_id=user.id,
                    farmer_id=farmer_id,
                    father_husband_name=father_name,
                    phone_number=str(phone),
                    gender=gender,
                    relation_type=relation_type,
                    house_no=str(house_no) if house_no else None,
                    street=street,
                    village=village,
                    mandal=mandal,
                    district=district,
                    state=state,
                    pincode=str(pincode),
                    total_area=float(total_area),
                    aadhaar_last_4=str(aadhaar_last_4),
                    bank_name=bank_name,
                    account_number=str(account_number),
                    ifsc_code=ifsc_code,
                )
                session.add(profile)
                await session.flush()

                if land_serial_1 and land_area_1:
                    session.add(LandRecord(
                        farmer_profile_id=profile.id,
                        serial_number=str(land_serial_1),
                        area=float(land_area_1),
                    ))
                if land_serial_2 and land_area_2:
                    session.add(LandRecord(
                        farmer_profile_id=profile.id,
                        serial_number=str(land_serial_2),
                        area=float(land_area_2),
                    ))

                farmer_idx_to_uid[int(farmer_idx)] = user.id

            await session.commit()

        pct = int((batch_end / total_farmers) * 100)
        print(f"    Farmers: {batch_end:,}/{total_farmers:,} ({pct}%)")

    print(f"  [OK] {total_farmers:,} farmers + profiles + land records inserted")

    # ─── 2. Generate Crop History for Farmers ─────────────────────────────────
    print(f"\n{'='*60}")
    print(f"  GENERATING CROP HISTORY (2-5 years per farmer)")
    print(f"{'='*60}")

    crop_count = expense_count = harvest_count = sale_count = 0
    farmer_uids = list(farmer_idx_to_uid.values())

    for batch_start in range(0, len(farmer_uids), batch_size):
        batch_end = min(batch_start + batch_size, len(farmer_uids))
        batch_ids = farmer_uids[batch_start:batch_end]

        async with async_session() as session:
            for user_id in batch_ids:
                years_back = random.randint(2, 5)
                farmer_area = random.uniform(1.5, 20.0)

                for year_offset in range(years_back):
                    year = now.year - year_offset
                    seasons_to_grow = random.sample(["Kharif", "Rabi"], k=random.choice([1, 2]))

                    for season in seasons_to_grow:
                        crop_data = random.choice(CROPS_DATA[season])
                        crop_area = round(random.uniform(0.5, min(farmer_area, 10.0)), 1)
                        variety = random.choice(crop_data["varieties"])

                        if season == "Kharif":
                            sowing_month = random.randint(6, 7)
                            harvest_month = random.randint(10, 12)
                        else:
                            sowing_month = random.randint(10, 11)
                            harvest_month = random.randint(2, 4)

                        sowing_date = datetime(year, sowing_month, random.randint(1, 28))
                        if season == "Rabi":
                            expected_harvest = datetime(year + 1 if harvest_month <= 4 else year, harvest_month, random.randint(1, 28))
                        else:
                            expected_harvest = datetime(year, harvest_month, random.randint(1, 28))

                        is_past = expected_harvest < now
                        status = "Harvested" if is_past else "Growing"
                        actual_harvest = expected_harvest + timedelta(days=random.randint(-5, 10)) if is_past else None

                        actual_yield = round(random.uniform(*crop_data["yield_range"]) * crop_area, 1) if is_past else 0
                        price_per_unit = round(random.uniform(*crop_data["price_range"]), 0) if is_past else 0
                        total_revenue = round(actual_yield * price_per_unit, 0) if is_past else 0
                        total_cost_val = round(random.uniform(*crop_data["cost_per_acre"]) * crop_area, 0)
                        net_profit = round(total_revenue - total_cost_val, 0) if is_past else 0

                        crop = Crop(
                            user_id=user_id,
                            name=crop_data["name"],
                            area=crop_area,
                            season=season,
                            variety=variety,
                            sowing_date=sowing_date,
                            expected_harvest_date=expected_harvest,
                            status=status,
                            crop_type=crop_data["type"],
                            actual_harvest_date=actual_harvest,
                            actual_yield=actual_yield,
                            selling_price_per_unit=price_per_unit,
                            total_revenue=total_revenue,
                            total_cost=total_cost_val,
                            net_profit=net_profit,
                            created_at=sowing_date,
                        )
                        session.add(crop)
                        await session.flush()
                        crop_count += 1

                        # Expenses
                        for _ in range(random.randint(3, 7)):
                            cat = random.choice(list(EXPENSE_TYPES.keys()))
                            exp_data = random.choice(EXPENSE_TYPES[cat])
                            exp_name, unit, qmin, qmax, ucmin, ucmax = exp_data
                            qty = round(random.uniform(qmin, qmax), 1)
                            unit_cost = round(random.uniform(ucmin, ucmax), 0)
                            exp_date = rand_date_between(sowing_date, actual_harvest if actual_harvest else expected_harvest)

                            session.add(CropExpense(
                                crop_id=crop.id, category=cat, type=exp_name,
                                quantity=qty, unit=unit, unit_cost=unit_cost,
                                total_cost=round(qty * unit_cost, 0),
                                date=exp_date,
                                payment_mode=random.choice(["cash", "digital"]),
                                unit_size=1.0,
                                duration=qty if cat == "Labor" else 1.0,
                                stage=random.choice(EXPENSE_STAGES),
                            ))
                            expense_count += 1

                        # Harvest & Sale for past crops
                        if is_past and actual_yield > 0:
                            for h_i in range(random.randint(1, 3)):
                                h_qty = round(actual_yield / random.randint(1, 3) * random.uniform(0.7, 1.3), 1)
                                h_date = actual_harvest + timedelta(days=h_i * random.randint(5, 15))
                                session.add(CropHarvest(
                                    crop_id=crop.id, date=h_date,
                                    stage=["First Picking", "Second Picking", "Final Harvest"][min(h_i, 2)],
                                    quantity=h_qty, unit="Quintals",
                                    quality=random.choice(["Grade A", "Grade A", "Grade B", "Grade C"]),
                                    selling_price_per_unit=price_per_unit,
                                    total_revenue=round(h_qty * price_per_unit, 0),
                                    buyer_type=random.choice(["Market", "Private", "Government", "Mill"]),
                                    sold_to=random.choice(["Local Mandi", "FCI", "Private Trader", "Nearby Mill", "Direct Buyer"]),
                                    status="Sold", created_at=h_date,
                                ))
                                harvest_count += 1

                            sale_date = actual_harvest + timedelta(days=random.randint(1, 30))
                            bag_size = random.choice([50, 75, 100])
                            total_bags = max(1, int(actual_yield * 100 / bag_size))
                            session.add(CropSale(
                                crop_id=crop.id, date=sale_date,
                                buyer_type=random.choice(["Mill", "Market", "Trader", "Direct"]),
                                buyer_name=random.choice(["Local Mandi", "Sri Lakshmi Mill", "FCI Depot", "Private Trader", "Agri Market Yard"]),
                                quantity_quintals=actual_yield, total_bags=total_bags, bag_size=bag_size,
                                price_per_quintal=price_per_unit, total_revenue=total_revenue,
                                payment_mode=random.choice(["cash", "digital", "cash"]),
                                status="sold", created_at=sale_date,
                            ))
                            sale_count += 1

            await session.commit()

        pct = int((batch_end / len(farmer_uids)) * 100)
        print(f"    Crop history: {batch_end:,}/{len(farmer_uids):,} ({pct}%) | "
              f"C:{crop_count:,} E:{expense_count:,} H:{harvest_count:,} S:{sale_count:,}")

    print(f"  [OK] {crop_count:,} crops, {expense_count:,} expenses, {harvest_count:,} harvests, {sale_count:,} sales")

    # ─── 3. Load Shops ────────────────────────────────────────────────────────
    ws = wb["shops"]
    shop_rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_shops = len(shop_rows)
    print(f"\n{'='*60}")
    print(f"  LOADING {total_shops:,} SHOPS")
    print(f"{'='*60}")

    async with async_session() as session:
        for row in shop_rows:
            (shop_idx, email, phone, password, full_name,
             shop_name, license_number, shop_id, father_name, relation_type, owner_name,
             aadhaar_number, pan_number,
             shop_address, landmark,
             house_no, street, village, mandal, district, state, pincode,
             region_name, region_idx, market_town,
             latitude, longitude,
             bank_name, account_number, ifsc_code,
             specialization) = row

            user = User(
                email=email, phone_number=str(phone), full_name=full_name,
                role=UserRole.SHOP, is_active=True, hashed_password=shop_hash,
            )
            session.add(user)
            await session.flush()

            profile = ShopProfile(
                user_id=user.id, shop_name=shop_name, license_number=license_number,
                shop_id=shop_id, father_name=father_name, relation_type=relation_type,
                owner_name=owner_name, aadhaar_number=str(aadhaar_number), pan_number=pan_number,
                shop_address=shop_address, landmark=landmark,
                house_no=str(house_no) if house_no else None,
                street=street, village=village, mandal=mandal, district=district,
                state=state, pincode=str(pincode),
                bank_name=bank_name, account_number=str(account_number), ifsc_code=ifsc_code,
            )
            session.add(profile)
            shop_idx_to_uid[int(shop_idx)] = user.id

        await session.commit()
    print(f"  [OK] {total_shops:,} shops + profiles inserted")

    # ─── 4. Generate Products & Shop Orders ───────────────────────────────────
    print(f"\n  Generating shop products & orders (using community links)...")
    product_count = order_count = 0

    # Read community links to build farmer→shop mapping
    ws_cl = wb["community_links"]
    cl_rows = list(ws_cl.iter_rows(min_row=2, values_only=True))

    # Build shop→farmer mappings from community_links
    shop_to_farmers = {}  # shop_idx → [farmer_idx, ...]
    for cl_row in cl_rows:
        farmer_idx_cl = cl_row[0]
        link_type = cl_row[15] if len(cl_row) > 15 else ""
        if link_type == "farmer_to_shop" and cl_row[7] is not None and cl_row[7] != "":
            s_idx = int(cl_row[7])
            shop_to_farmers.setdefault(s_idx, []).append(int(farmer_idx_cl))

    for batch_start in range(0, len(shop_idx_to_uid), 50):
        batch_end = min(batch_start + 50, len(shop_idx_to_uid))
        batch_items = list(shop_idx_to_uid.items())[batch_start:batch_end]

        async with async_session() as session:
            for s_idx, shop_uid in batch_items:
                num_products = random.randint(8, 15)
                selected_products = random.sample(SHOP_PRODUCTS_DATA, min(num_products, len(SHOP_PRODUCTS_DATA)))
                product_ids_for_shop = []

                for pd in selected_products:
                    qty = random.randint(20, 500)
                    product = Product(
                        user_id=shop_uid, name=pd["name"], category=pd["category"],
                        brand=pd["brand"], price=pd["price"], cost_price=pd["cost"],
                        quantity=qty, unit=pd["unit"], quantity_per_unit=pd["qty_per_unit"],
                        measure_unit="kg", batch_number=f"BATCH-{random.randint(10000, 99999)}",
                        description=f"{pd['name']} - {pd['brand']}",
                        low_stock_threshold=10, status="active",
                        created_at=rand_date_between(datetime(2022, 1, 1), now),
                    )
                    session.add(product)
                    await session.flush()
                    product_ids_for_shop.append((product.id, pd))
                    product_count += 1

                # Generate orders — use linked farmers from community
                linked_farmer_idxs = shop_to_farmers.get(s_idx, [])
                num_orders = random.randint(30, 120)

                for _ in range(num_orders):
                    order_date = rand_date_between(datetime(2023, 1, 1), now)
                    num_items = random.randint(1, 4)
                    items_selected = random.sample(product_ids_for_shop, min(num_items, len(product_ids_for_shop)))

                    total = 0.0
                    order_items = []
                    for pid, pd in items_selected:
                        qty = random.randint(1, 10)
                        subtotal = round(pd["price"] * qty, 0)
                        total += subtotal
                        order_items.append((pid, pd["name"], qty, pd["price"], subtotal))

                    discount = round(total * random.choice([0, 0, 0, 0.02, 0.05, 0.1]), 0)
                    final_amount = total - discount

                    # 70% of orders from linked community farmers, 30% walk-in
                    if linked_farmer_idxs and random.random() > 0.3:
                        f_idx = random.choice(linked_farmer_idxs)
                        farmer_uid = farmer_idx_to_uid.get(f_idx)
                        farmer_name = f"Community Farmer #{f_idx}"
                    else:
                        farmer_uid = None
                        farmer_name = "Walk-in Customer"

                    order = ShopOrder(
                        shop_id=shop_uid, farmer_id=farmer_uid,
                        farmer_name=farmer_name,
                        total_amount=total, discount=discount, final_amount=final_amount,
                        payment_mode=random.choice(["cash", "cash", "upi", "credit"]),
                        payment_status="paid", status="completed",
                        created_at=order_date,
                    )
                    session.add(order)
                    await session.flush()
                    order_count += 1

                    for pid, pname, qty, price, subtotal in order_items:
                        session.add(ShopOrderItem(
                            order_id=order.id, product_id=pid,
                            product_name=pname, quantity=qty,
                            unit_price=price, subtotal=subtotal,
                        ))

                    if random.random() > 0.6:
                        session.add(ShopExpense(
                            order_id=order.id,
                            transportation=round(random.uniform(50, 500), 0),
                            labour=round(random.uniform(100, 800), 0),
                            other=round(random.uniform(0, 200), 0),
                            created_at=order_date,
                        ))

            await session.commit()
        print(f"    Shop products & orders: {batch_end}/{len(shop_idx_to_uid)} | P:{product_count:,} O:{order_count:,}")

    print(f"  [OK] {product_count:,} products, {order_count:,} orders")

    # ─── 5. Load Mills ────────────────────────────────────────────────────────
    ws = wb["mills"]
    mill_rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_mills = len(mill_rows)
    print(f"\n{'='*60}")
    print(f"  LOADING {total_mills:,} MILLS")
    print(f"{'='*60}")

    async with async_session() as session:
        for row in mill_rows:
            (mill_idx, email, phone, password, full_name,
             mill_name, license_number, mill_id, father_name, relation_type, owner_name,
             aadhaar_number, pan_number,
             house_no, street, village, mandal, district, state, pincode,
             region_name, region_idx, mill_area,
             latitude, longitude, location_text,
             bank_name, account_number, ifsc_code,
             processing_type, daily_capacity) = row

            user = User(
                email=email, phone_number=str(phone), full_name=full_name,
                role=UserRole.MANUFACTURER, is_active=True, hashed_password=mill_hash,
            )
            session.add(user)
            await session.flush()

            profile = MillProfile(
                user_id=user.id, mill_name=mill_name, license_number=license_number,
                mill_id=mill_id, father_name=father_name, relation_type=relation_type,
                owner_name=owner_name, aadhaar_number=str(aadhaar_number), pan_number=pan_number,
                house_no=str(house_no) if house_no else None,
                street=street, village=village, mandal=mandal, district=district,
                state=state, pincode=str(pincode), location_text=location_text,
                bank_name=bank_name, account_number=str(account_number), ifsc_code=ifsc_code,
            )
            session.add(profile)
            mill_idx_to_uid[int(mill_idx)] = user.id

        await session.commit()
    print(f"  [OK] {total_mills:,} mills + profiles inserted")

    # ─── 6. Generate Mill Purchases from Community Farmers ────────────────────
    print(f"\n  Generating mill purchases from community-linked farmers...")
    purchase_count = mill_expense_count = 0

    # Build mill→farmer mapping from community_links
    mill_to_farmers = {}
    for cl_row in cl_rows:
        farmer_idx_cl = cl_row[0]
        link_type = cl_row[15] if len(cl_row) > 15 else ""
        if link_type == "farmer_to_mill" and cl_row[11] is not None and cl_row[11] != "":
            m_idx = int(cl_row[11])
            mill_to_farmers.setdefault(m_idx, []).append(int(farmer_idx_cl))

    async with async_session() as session:
        for m_idx, mill_uid in mill_idx_to_uid.items():
            linked_farmers = mill_to_farmers.get(m_idx, [])
            num_purchases = random.randint(15, 40)

            for _ in range(num_purchases):
                p_date = rand_date_between(datetime(2023, 1, 1), now)
                crop_name = random.choice(["Paddy", "Wheat", "Maize", "Groundnut", "Soybean", "Sugarcane", "Chilli", "Turmeric"])
                qty = round(random.uniform(10, 200), 0)
                price_per = round(random.uniform(1500, 5000), 0)
                total_cost = round(qty * price_per, 0)
                transport = round(random.uniform(500, 5000), 0)

                # Use community-linked farmer if available
                if linked_farmers and random.random() > 0.2:
                    f_idx = random.choice(linked_farmers)
                    farmer_uid = farmer_idx_to_uid.get(f_idx)
                    farmer_name = f"Community Farmer #{f_idx}"
                else:
                    farmer_uid = random.choice(farmer_uids)
                    farmer_name = f"Farmer #{farmer_uid}"

                session.add(ManufacturerPurchase(
                    manufacturer_id=mill_uid, farmer_id=farmer_uid,
                    farmer_name=farmer_name, crop_name=crop_name,
                    quantity=qty, unit="quintals", price_per_unit=price_per,
                    total_cost=total_cost, transport_cost=transport,
                    quality_grade=random.choice(["A", "A", "B", "B", "C"]),
                    batch_id=f"M-PUR-{purchase_count + 1}",
                    date=p_date,
                ))
                purchase_count += 1

            # Mill expenses
            for _ in range(random.randint(8, 20)):
                session.add(ManufacturerExpense(
                    manufacturer_id=mill_uid,
                    category=random.choice(["electricity", "labour", "maintenance", "fuel", "rent", "packaging"]),
                    amount=round(random.uniform(5000, 100000), 0),
                    description=random.choice([
                        "Monthly electricity", "Workers salary", "Machine maintenance",
                        "Diesel for generator", "Warehouse rent", "Gunny bags purchase",
                        "Machine parts", "Water supply", "Security charges",
                    ]),
                    expense_date=rand_date_between(date(2023, 1, 1), date.today()),
                    created_at=datetime.utcnow(),
                ))
                mill_expense_count += 1

        await session.commit()
    print(f"  [OK] {purchase_count:,} purchases, {mill_expense_count:,} mill expenses")

    # ─── 7. Load Customers ────────────────────────────────────────────────────
    ws = wb["customers"]
    customer_rows = list(ws.iter_rows(min_row=2, values_only=True))
    total_customers = len(customer_rows)
    print(f"\n{'='*60}")
    print(f"  LOADING {total_customers:,} CUSTOMERS")
    print(f"{'='*60}")

    for batch_start in range(0, total_customers, batch_size):
        batch_end = min(batch_start + batch_size, total_customers)
        batch = customer_rows[batch_start:batch_end]

        async with async_session() as session:
            for row in batch:
                (cust_idx, email, phone, password, full_name,
                 father_name, relation_type, id_number,
                 house_no, street, village, mandal, district, state, pincode,
                 customer_region, preferred_products,
                 bank_name, account_number, ifsc_code) = row

                user = User(
                    email=email, phone_number=str(phone), full_name=full_name,
                    role=UserRole.CUSTOMER, is_active=True, hashed_password=customer_hash,
                )
                session.add(user)
                await session.flush()

                profile = CustomerProfile(
                    user_id=user.id, father_name=father_name,
                    phone_number=str(phone), relation_type=relation_type,
                    id_number=str(id_number),
                    house_no=str(house_no) if house_no else None,
                    street=street, village=village, mandal=mandal, district=district,
                    state=state, pincode=str(pincode),
                    bank_name=bank_name, account_number=str(account_number), ifsc_code=ifsc_code,
                )
                session.add(profile)
                customer_idx_to_uid[int(cust_idx)] = user.id

            await session.commit()
        pct = int((batch_end / total_customers) * 100)
        print(f"    Customers: {batch_end:,}/{total_customers:,} ({pct}%)")

    print(f"  [OK] {total_customers:,} customers + profiles inserted")

    # ─── 8. Generate Customer Orders (cross-region buying) ────────────────────
    print(f"\n  Generating customer orders (cross-region from farmer/mill)...")
    cust_order_count = 0

    # Read customer_orders sheet for cross-region order data
    ws_co = wb["customer_orders"]
    co_rows = list(ws_co.iter_rows(min_row=2, values_only=True))

    for batch_start in range(0, len(co_rows), batch_size):
        batch_end = min(batch_start + batch_size, len(co_rows))
        batch = co_rows[batch_start:batch_end]

        async with async_session() as session:
            for row in batch:
                (order_id, cust_idx, cust_name, cust_state, cust_district,
                 seller_type, seller_idx, seller_name, seller_region,
                 product_name, quantity, unit_price, total_amount,
                 order_status, payment_mode, order_date) = row

                cust_uid = customer_idx_to_uid.get(int(cust_idx))
                if not cust_uid:
                    continue

                if seller_type == "farmer":
                    seller_uid = farmer_idx_to_uid.get(int(seller_idx))
                elif seller_type == "mill":
                    seller_uid = mill_idx_to_uid.get(int(seller_idx))
                else:
                    seller_uid = None

                if not seller_uid:
                    seller_uid = random.choice(farmer_uids)

                order = CustomerOrder(
                    customer_id=cust_uid,
                    total_amount=float(total_amount),
                    status=order_status or "delivered",
                    created_at=datetime.strptime(str(order_date), "%Y-%m-%d") if order_date else now,
                )
                session.add(order)
                await session.flush()

                session.add(CustomerOrderItem(
                    order_id=order.id,
                    product_id=1,  # Placeholder — in real data this would link to actual product
                    seller_id=seller_uid,
                    product_name=product_name or "Agricultural Product",
                    quantity=float(quantity) if quantity else 1.0,
                    price=float(unit_price) if unit_price else 0.0,
                ))
                cust_order_count += 1

            await session.commit()
        print(f"    Customer orders: {min(batch_end, len(co_rows)):,}/{len(co_rows):,}")

    print(f"  [OK] {cust_order_count:,} customer orders inserted")

    # ─── Final Summary ────────────────────────────────────────────────────────
    print(f"\n{'='*60}")
    print(f"  SEED DATA V2 LOAD COMPLETE")
    print(f"{'='*60}")
    print(f"  Users:                {total_farmers + total_shops + total_mills + total_customers:>10,}")
    print(f"    Farmers:            {total_farmers:>10,}")
    print(f"    Shops:              {total_shops:>10,}")
    print(f"    Mills:              {total_mills:>10,}")
    print(f"    Customers:          {total_customers:>10,}")
    print(f"  Crops:                {crop_count:>10,}")
    print(f"  Crop Expenses:        {expense_count:>10,}")
    print(f"  Harvests:             {harvest_count:>10,}")
    print(f"  Crop Sales:           {sale_count:>10,}")
    print(f"  Products:             {product_count:>10,}")
    print(f"  Shop Orders:          {order_count:>10,}")
    print(f"  Mill Purchases:       {purchase_count:>10,}")
    print(f"  Mill Expenses:        {mill_expense_count:>10,}")
    print(f"  Customer Orders:      {cust_order_count:>10,}")
    print(f"{'='*60}")
    print(f"\n  Passwords:")
    print(f"    Farmer:   Farmer@123")
    print(f"    Shop:     Shop@123")
    print(f"    Mill:     Mill@123")
    print(f"    Customer: Customer@123")


if __name__ == "__main__":
    asyncio.run(load_all())
