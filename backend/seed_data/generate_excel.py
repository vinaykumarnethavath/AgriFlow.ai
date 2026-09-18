"""
AgriFlow Seed Data Generator
=============================
Generates a comprehensive Excel workbook with 4 sheets:
  1. farmers   – 10,000 farmers across 10 Indian states
  2. shops     – 500 agri-input shops
  3. mills     – 200 rice/flour mills
  4. customers – 2,000 end-consumers

Run:  python generate_excel.py
Output: agriflow_seed_data.xlsx (same directory)
"""

import random
import string
import os
from datetime import datetime, timedelta, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ─── Deterministic seed for reproducibility ──────────────────────────────────
random.seed(42)

# ─── Location Master Data ────────────────────────────────────────────────────

LOCATIONS = [
    # (state, [(district, mandal, village, pincode_prefix), ...])
    ("Telangana", [
        ("Warangal", "Hanamkonda", ["Kazipet", "Elkathurthy", "Parkal", "Narsampet", "Jangaon"], "506"),
        ("Karimnagar", "Karimnagar", ["Huzurabad", "Jagtial", "Metpally", "Korutla", "Peddapalli"], "505"),
        ("Nizamabad", "Nizamabad", ["Bodhan", "Armoor", "Kamareddy", "Banswada", "Yellareddy"], "503"),
        ("Nalgonda", "Nalgonda", ["Miryalaguda", "Suryapet", "Devarakonda", "Kodad", "Nakrekal"], "508"),
        ("Khammam", "Khammam", ["Kothagudem", "Bhadrachalam", "Sathupalli", "Yellandu", "Madhira"], "507"),
    ]),
    ("Andhra Pradesh", [
        ("Guntur", "Guntur", ["Tenali", "Mangalagiri", "Sattenapalli", "Narasaraopet", "Bapatla"], "522"),
        ("Krishna", "Machilipatnam", ["Vijayawada", "Gudivada", "Nuzvid", "Jaggayyapeta", "Nandigama"], "521"),
        ("East Godavari", "Kakinada", ["Rajahmundry", "Amalapuram", "Mandapeta", "Ramachandrapuram", "Pithapuram"], "533"),
        ("West Godavari", "Eluru", ["Bhimavaram", "Narsapuram", "Tadepalligudem", "Palacole", "Tanuku"], "534"),
        ("Kurnool", "Kurnool", ["Nandyal", "Adoni", "Dhone", "Yemmiganur", "Allagadda"], "518"),
    ]),
    ("Uttar Pradesh", [
        ("Lucknow", "Lucknow", ["Mohanlalganj", "Malihabad", "Bakshi Ka Talab", "Chinhat", "Kakori"], "226"),
        ("Varanasi", "Varanasi", ["Pindra", "Chiraigaon", "Rajatalab", "Cholapur", "Harahua"], "221"),
        ("Agra", "Agra", ["Fatehabad", "Kheragarh", "Etmadpur", "Bah", "Shamshabad"], "282"),
        ("Prayagraj", "Prayagraj", ["Phulpur", "Soraon", "Handia", "Meja", "Koraon"], "211"),
        ("Kanpur", "Kanpur Nagar", ["Bilhaur", "Ghatampur", "Bhitargaon", "Sarsaul", "Patara"], "208"),
    ]),
    ("Maharashtra", [
        ("Pune", "Haveli", ["Baramati", "Indapur", "Daund", "Junnar", "Shirur"], "411"),
        ("Nashik", "Nashik", ["Sinnar", "Niphad", "Dindori", "Igatpuri", "Malegaon"], "422"),
        ("Nagpur", "Nagpur", ["Ramtek", "Kamptee", "Umred", "Katol", "Saoner"], "440"),
        ("Chhatrapati Sambhajinagar", "Aurangabad", ["Paithan", "Kannad", "Khuldabad", "Gangapur", "Vaijapur"], "431"),
        ("Kolhapur", "Kolhapur", ["Ichalkaranji", "Hatkanangale", "Shirol", "Kagal", "Radhanagari"], "416"),
    ]),
    ("Karnataka", [
        ("Mysuru", "Mysuru", ["Nanjangud", "T Narasipura", "Hunsur", "Periyapatna", "K R Nagar"], "570"),
        ("Belgaum", "Belgaum", ["Gokak", "Chikkodi", "Athani", "Raibag", "Ramdurg"], "591"),
        ("Dharwad", "Dharwad", ["Hubli", "Kundgol", "Navalgund", "Kalghatgi", "Annigeri"], "580"),
        ("Hassan", "Hassan", ["Belur", "Channarayapatna", "Arkalgud", "Arsikere", "Holenarasipura"], "573"),
        ("Mandya", "Mandya", ["Maddur", "Malavalli", "Srirangapatna", "Pandavapura", "Nagamangala"], "571"),
    ]),
    ("Tamil Nadu", [
        ("Thanjavur", "Thanjavur", ["Kumbakonam", "Pattukkottai", "Orathanadu", "Thiruvaiyaru", "Papanasam"], "613"),
        ("Coimbatore", "Coimbatore", ["Pollachi", "Mettupalayam", "Annur", "Sulur", "Kinathukadavu"], "641"),
        ("Salem", "Salem", ["Attur", "Mettur", "Omalur", "Edappadi", "Sankagiri"], "636"),
        ("Madurai", "Madurai", ["Melur", "Vadipatti", "Usilampatti", "Thirumangalam", "Peraiyur"], "625"),
        ("Tiruchirapalli", "Tiruchirapalli", ["Musiri", "Lalgudi", "Thuraiyur", "Manapparai", "Srirangam"], "620"),
    ]),
    ("Punjab", [
        ("Ludhiana", "Ludhiana", ["Jagraon", "Samrala", "Khanna", "Raikot", "Payal"], "141"),
        ("Amritsar", "Amritsar", ["Tarn Taran", "Ajnala", "Majitha", "Jandiala Guru", "Baba Bakala"], "143"),
        ("Patiala", "Patiala", ["Rajpura", "Nabha", "Samana", "Patran", "Ghanaur"], "147"),
        ("Jalandhar", "Jalandhar", ["Nakodar", "Phillaur", "Shahkot", "Adampur", "Kartarpur"], "144"),
        ("Bathinda", "Bathinda", ["Rampura Phul", "Talwandi Sabo", "Maur", "Goniana", "Sangat"], "151"),
    ]),
    ("Madhya Pradesh", [
        ("Indore", "Indore", ["Mhow", "Depalpur", "Sanwer", "Hatod", "Betma"], "452"),
        ("Bhopal", "Bhopal", ["Berasia", "Sehore", "Ashta", "Budni", "Ichhawar"], "462"),
        ("Jabalpur", "Jabalpur", ["Sihora", "Katni", "Patan", "Shahpura", "Bargi"], "482"),
        ("Gwalior", "Gwalior", ["Dabra", "Bhitarwar", "Morar", "Pichhore", "Hastinapur"], "474"),
        ("Ujjain", "Ujjain", ["Nagda", "Tarana", "Mahidpur", "Barnagar", "Ghattia"], "456"),
    ]),
    ("Gujarat", [
        ("Ahmedabad", "Ahmedabad", ["Dholka", "Dhandhuka", "Bavla", "Sanand", "Viramgam"], "380"),
        ("Rajkot", "Rajkot", ["Gondal", "Jetpur", "Dhoraji", "Upleta", "Jasdan"], "360"),
        ("Surat", "Surat", ["Bardoli", "Kamrej", "Mandvi", "Olpad", "Chorasi"], "394"),
        ("Vadodara", "Vadodara", ["Savli", "Padra", "Karjan", "Dabhoi", "Waghodia"], "390"),
        ("Junagadh", "Junagadh", ["Visavadar", "Keshod", "Mangrol", "Vanthali", "Mendarda"], "362"),
    ]),
    ("West Bengal", [
        ("Purba Bardhaman", "Bardhaman", ["Memari", "Kalna", "Katwa", "Mongalkote", "Ausgram"], "713"),
        ("Hooghly", "Chinsurah", ["Arambagh", "Chandannagar", "Serampore", "Pursurah", "Goghat"], "712"),
        ("Nadia", "Krishnanagar", ["Ranaghat", "Santipur", "Nabadwip", "Kalyani", "Tehatta"], "741"),
        ("Murshidabad", "Baharampur", ["Lalbag", "Kandi", "Domkal", "Jalangi", "Raghunathganj"], "742"),
        ("Malda", "English Bazar", ["Gazole", "Habibpur", "Old Malda", "Manikchak", "Ratua"], "732"),
    ]),
]

# ─── Indian Names ─────────────────────────────────────────────────────────────

MALE_FIRST_NAMES = [
    "Ramesh", "Suresh", "Mahesh", "Rajesh", "Ganesh", "Venkatesh", "Naresh", "Dinesh",
    "Anil", "Sunil", "Manoj", "Ajay", "Vijay", "Sanjay", "Ravi", "Kiran",
    "Prakash", "Harish", "Satish", "Girish", "Ashok", "Vinod", "Pramod", "Arvind",
    "Deepak", "Alok", "Pankaj", "Mukesh", "Yogesh", "Umesh", "Rakesh", "Nilesh",
    "Gopal", "Mohan", "Sohan", "Rohan", "Vishal", "Nikhil", "Sachin", "Rahul",
    "Amit", "Sumit", "Mohit", "Rohit", "Varun", "Tarun", "Charan", "Karan",
    "Naveen", "Praveen", "Sravan", "Pavan", "Bharat", "Ishwar", "Shyam", "Ram",
    "Krishna", "Balaji", "Murali", "Srinivas", "Venu", "Prasad", "Lakshman", "Raju",
    "Sathya", "Bhaskar", "Shankar", "Sekhar", "Kishore", "Bala", "Anand", "Nanda",
    "Jagdish", "Devendra", "Surendra", "Narendra", "Rajendra", "Mahendra", "Birendra", "Dhirendra",
    "Santosh", "Ramakrishna", "Venkata", "Siddharth", "Akhil", "Pranav", "Arjun", "Vikram",
    "Mahadev", "Shiva", "Ganpat", "Hanuman", "Tulsi", "Govind", "Hari", "Nagendra",
    "Basavaraj", "Channappa", "Yellappa", "Mallappa", "Thimmappa", "Siddappa", "Hanumanth", "Basappa",
]

FEMALE_FIRST_NAMES = [
    "Lakshmi", "Saraswati", "Parvati", "Durga", "Sita", "Radha", "Ganga", "Yamuna",
    "Anita", "Sunita", "Kavita", "Savita", "Mamta", "Sumitra", "Rekha", "Meena",
    "Padma", "Kamla", "Shanti", "Jaya", "Vijaya", "Sujata", "Renuka", "Anuradha",
    "Bhavani", "Gayatri", "Lata", "Geeta", "Seema", "Neelam", "Pushpa", "Saroja",
    "Vasanta", "Chandra", "Indira", "Nirmala", "Usha", "Aruna", "Swati", "Priya",
]

LAST_NAMES = [
    "Reddy", "Naidu", "Rao", "Kumar", "Singh", "Yadav", "Sharma", "Verma",
    "Patel", "Shah", "Desai", "Patil", "Shinde", "Jadhav", "More", "Gowda",
    "Swamy", "Nair", "Menon", "Pillai", "Das", "Ghosh", "Roy", "Sen",
    "Joshi", "Kulkarni", "Deshpande", "Hegde", "Shetty", "Bhat", "Kaur", "Gill",
    "Sidhu", "Dhillon", "Sandhu", "Chauhan", "Thakur", "Rajput", "Tiwari", "Pandey",
    "Mishra", "Dubey", "Gupta", "Agarwal", "Bansal", "Jain", "Sethi", "Mehta",
    "Chaudhary", "Mandal", "Sarkar", "Mondal", "Biswas", "Barman", "Mahato", "Murmu",
    "Nethavath", "Rathod", "Jadhav", "Pawar", "Bhosale", "Chavan", "Kale", "Solanki",
]

FATHER_NAMES = [
    "Ramaiah", "Venkataiah", "Narasimha", "Balaiah", "Chandraiah", "Laxmaiah", "Mallaiah",
    "Rajanna", "Pochaiah", "Yellaiah", "Veeraiah", "Bhaskar", "Gangaram", "Laxman",
    "Bhagwan", "Dharma", "Shankar", "Kishan", "Mohan", "Gopal", "Shyam", "Girdhari",
    "Raghunath", "Jagannath", "Vishwanath", "Dashrath", "Trilok", "Omkar", "Devidas",
]

# ─── Crop Master Data ─────────────────────────────────────────────────────────

CROPS = {
    "Kharif": [
        {"name": "Paddy", "type": "Cereal", "varieties": ["Sona Masuri", "BPT 5204", "IR-64", "Swarna", "Tellahamsa", "MTU-1010"], "yield_range": (15, 35), "price_range": (1800, 2500), "cost_per_acre": (12000, 22000)},
        {"name": "Maize", "type": "Cereal", "varieties": ["DHM-117", "Kaveri 50", "NK-6240", "Pioneer 3522"], "yield_range": (20, 40), "price_range": (1400, 2100), "cost_per_acre": (10000, 18000)},
        {"name": "Cotton", "type": "Commercial", "varieties": ["Bt Cotton", "Bunny", "Mallika", "Suraj", "NCS-145"], "yield_range": (8, 18), "price_range": (5500, 7200), "cost_per_acre": (15000, 30000)},
        {"name": "Soybean", "type": "Oilseed", "varieties": ["JS-335", "JS-9560", "NRC-37", "MACS-1407"], "yield_range": (8, 15), "price_range": (3800, 5200), "cost_per_acre": (8000, 15000)},
        {"name": "Groundnut", "type": "Oilseed", "varieties": ["TMV-2", "JL-501", "Kadiri-6", "TAG-24"], "yield_range": (10, 20), "price_range": (4500, 6500), "cost_per_acre": (12000, 20000)},
        {"name": "Chilli", "type": "Spice", "varieties": ["Teja", "Byadgi", "S-4", "LCA-334", "US-341"], "yield_range": (5, 15), "price_range": (8000, 18000), "cost_per_acre": (25000, 45000)},
        {"name": "Turmeric", "type": "Spice", "varieties": ["Erode Local", "Salem", "Rajapore", "Mydukur"], "yield_range": (20, 40), "price_range": (6000, 12000), "cost_per_acre": (30000, 55000)},
        {"name": "Sugarcane", "type": "Commercial", "varieties": ["Co-86032", "CoC-671", "Co-0238", "CoM-0265"], "yield_range": (300, 500), "price_range": (280, 350), "cost_per_acre": (35000, 60000)},
        {"name": "Jowar", "type": "Cereal", "varieties": ["CSV-15", "CSH-16", "Maldandi", "M-35-1"], "yield_range": (8, 18), "price_range": (2200, 3200), "cost_per_acre": (6000, 12000)},
        {"name": "Bajra", "type": "Cereal", "varieties": ["HHB-67", "ICTP-8203", "Raj-171", "GHB-558"], "yield_range": (8, 16), "price_range": (1900, 2800), "cost_per_acre": (5000, 10000)},
    ],
    "Rabi": [
        {"name": "Wheat", "type": "Cereal", "varieties": ["PBW-343", "HD-2967", "Lok-1", "GW-496", "WH-1105"], "yield_range": (15, 30), "price_range": (1900, 2600), "cost_per_acre": (10000, 18000)},
        {"name": "Chickpea", "type": "Pulse", "varieties": ["JG-11", "JAKI-9218", "Vijay", "Vishal", "KAK-2"], "yield_range": (6, 12), "price_range": (4200, 5800), "cost_per_acre": (8000, 14000)},
        {"name": "Mustard", "type": "Oilseed", "varieties": ["Pusa Bold", "RH-749", "Bio-902", "NRCHB-101"], "yield_range": (6, 12), "price_range": (4500, 6000), "cost_per_acre": (6000, 11000)},
        {"name": "Onion", "type": "Vegetable", "varieties": ["Nasik Red", "Bellary Red", "Pusa Ratnar", "Agrifound Dark Red"], "yield_range": (80, 150), "price_range": (800, 2500), "cost_per_acre": (30000, 50000)},
        {"name": "Potato", "type": "Vegetable", "varieties": ["Kufri Jyoti", "Kufri Pukhraj", "Kufri Bahar", "Kufri Chipsona"], "yield_range": (80, 150), "price_range": (600, 1500), "cost_per_acre": (35000, 55000)},
        {"name": "Sunflower", "type": "Oilseed", "varieties": ["KBSH-44", "MSFH-17", "BSH-1", "Sunbred-275"], "yield_range": (5, 10), "price_range": (4000, 5500), "cost_per_acre": (7000, 12000)},
        {"name": "Bengal Gram", "type": "Pulse", "varieties": ["JG-11", "JG-14", "KAK-2", "Vijay"], "yield_range": (5, 10), "price_range": (4500, 6200), "cost_per_acre": (7000, 12000)},
        {"name": "Lentil", "type": "Pulse", "varieties": ["IPL-316", "K-75", "Pant L-5", "HUL-57"], "yield_range": (4, 8), "price_range": (3800, 5500), "cost_per_acre": (5000, 9000)},
    ],
    "Zaid": [
        {"name": "Watermelon", "type": "Fruit", "varieties": ["Sugar Baby", "Arka Manik", "Crimson Sweet"], "yield_range": (100, 200), "price_range": (400, 1000), "cost_per_acre": (20000, 35000)},
        {"name": "Muskmelon", "type": "Fruit", "varieties": ["Pusa Sharbati", "Arka Rajhans", "Punjab Sunehri"], "yield_range": (60, 120), "price_range": (600, 1500), "cost_per_acre": (18000, 30000)},
        {"name": "Cucumber", "type": "Vegetable", "varieties": ["Pusa Uday", "Swarna Sheetal", "Malini"], "yield_range": (50, 100), "price_range": (500, 1200), "cost_per_acre": (15000, 25000)},
        {"name": "Green Gram", "type": "Pulse", "varieties": ["IPM-02-3", "SML-668", "Pusa Vishal"], "yield_range": (3, 7), "price_range": (5000, 7200), "cost_per_acre": (5000, 9000)},
    ],
}

EXPENSE_CATEGORIES = {
    "Input": ["Seed", "Fertilizer (DAP)", "Fertilizer (Urea)", "Fertilizer (MOP)", "Fertilizer (NPK)", "Pesticide", "Herbicide", "Fungicide", "Growth Regulator", "Micronutrients"],
    "Labor": ["Ploughing Labour", "Sowing Labour", "Weeding Labour", "Spraying Labour", "Harvesting Labour", "Loading/Unloading"],
    "Machinery": ["Tractor Hire", "Rotavator", "Harvester Hire", "Sprayer Rent", "Thresher"],
    "Irrigation": ["Borewell Electricity", "Diesel Pump Fuel", "Drip Maintenance", "Canal Water Charge"],
    "Logistics": ["Transport to Mandi", "Transport to Mill", "Gunny Bags", "Storage Rent"],
}

EXPENSE_STAGES = ["Sowing", "Germination", "Vegetative", "Flowering", "Fruiting", "Harvesting", "Post-Harvest"]

# ─── Bank Data ────────────────────────────────────────────────────────────────

BANKS = [
    ("State Bank of India", "SBIN"),
    ("Andhra Bank", "ANDB"),
    ("Canara Bank", "CNRB"),
    ("Bank of Baroda", "BARB"),
    ("Punjab National Bank", "PUNB"),
    ("Union Bank of India", "UBIN"),
    ("Indian Bank", "IDIB"),
    ("Central Bank of India", "CBIN"),
    ("Bank of India", "BKID"),
    ("HDFC Bank", "HDFC"),
    ("ICICI Bank", "ICIC"),
    ("Axis Bank", "UTIB"),
    ("Kotak Mahindra Bank", "KKBK"),
    ("IndusInd Bank", "INDB"),
    ("Telangana Grameena Bank", "TGMB"),
    ("Andhra Pradesh Grameena Vikas Bank", "APGV"),
    ("Syndicate Bank", "SYNB"),
]

# ─── Shop Products ───────────────────────────────────────────────────────────

SHOP_PRODUCTS = [
    {"name": "DAP Fertilizer", "category": "fertilizer", "brand": "IFFCO", "unit": "bag", "qty_per_unit": 50, "cost": 1150, "price": 1350},
    {"name": "Urea", "category": "fertilizer", "brand": "IFFCO", "unit": "bag", "qty_per_unit": 50, "cost": 266, "price": 300},
    {"name": "MOP Potash", "category": "fertilizer", "brand": "IPL", "unit": "bag", "qty_per_unit": 50, "cost": 850, "price": 1000},
    {"name": "NPK 20-20-0", "category": "fertilizer", "brand": "Coromandel", "unit": "bag", "qty_per_unit": 50, "cost": 1100, "price": 1280},
    {"name": "NPK 10-26-26", "category": "fertilizer", "brand": "Zuari", "unit": "bag", "qty_per_unit": 50, "cost": 1200, "price": 1400},
    {"name": "Zinc Sulphate", "category": "fertilizer", "brand": "Tata", "unit": "kg", "qty_per_unit": 25, "cost": 40, "price": 55},
    {"name": "Gypsum", "category": "fertilizer", "brand": "GSFC", "unit": "bag", "qty_per_unit": 50, "cost": 200, "price": 280},
    {"name": "SSP", "category": "fertilizer", "brand": "Paradeep", "unit": "bag", "qty_per_unit": 50, "cost": 350, "price": 450},
    {"name": "Imidacloprid 17.8 SL", "category": "pesticide", "brand": "Bayer", "unit": "liter", "qty_per_unit": 1, "cost": 800, "price": 1050},
    {"name": "Chlorpyriphos 20 EC", "category": "pesticide", "brand": "Dhanuka", "unit": "liter", "qty_per_unit": 1, "cost": 350, "price": 480},
    {"name": "Monocrotophos 36 SL", "category": "pesticide", "brand": "Syngenta", "unit": "liter", "qty_per_unit": 1, "cost": 300, "price": 420},
    {"name": "Carbendazim 50 WP", "category": "pesticide", "brand": "BASF", "unit": "packet", "qty_per_unit": 0.5, "cost": 180, "price": 260},
    {"name": "Mancozeb 75 WP", "category": "pesticide", "brand": "UPL", "unit": "packet", "qty_per_unit": 1, "cost": 250, "price": 350},
    {"name": "Glyphosate 41 SL", "category": "pesticide", "brand": "Excel", "unit": "liter", "qty_per_unit": 1, "cost": 350, "price": 500},
    {"name": "Neem Oil", "category": "pesticide", "brand": "Multiplex", "unit": "liter", "qty_per_unit": 1, "cost": 200, "price": 300},
    {"name": "Paddy Seed BPT-5204", "category": "seeds", "brand": "NSC", "unit": "kg", "qty_per_unit": 10, "cost": 45, "price": 65},
    {"name": "Cotton Seed Bt", "category": "seeds", "brand": "Mahyco", "unit": "packet", "qty_per_unit": 0.45, "cost": 650, "price": 850},
    {"name": "Maize Seed Hybrid", "category": "seeds", "brand": "Kaveri", "unit": "kg", "qty_per_unit": 5, "cost": 180, "price": 250},
    {"name": "Wheat Seed HD-2967", "category": "seeds", "brand": "ICAR", "unit": "kg", "qty_per_unit": 40, "cost": 30, "price": 45},
    {"name": "Drip Irrigation Kit", "category": "equipment", "brand": "Jain", "unit": "set", "qty_per_unit": 1, "cost": 8000, "price": 12000},
    {"name": "Sprayer Manual 16L", "category": "equipment", "brand": "Neptune", "unit": "piece", "qty_per_unit": 1, "cost": 900, "price": 1400},
    {"name": "Battery Sprayer 16L", "category": "equipment", "brand": "Aspee", "unit": "piece", "qty_per_unit": 1, "cost": 2500, "price": 3800},
    {"name": "Vermicompost", "category": "fertilizer", "brand": "Organic Gold", "unit": "bag", "qty_per_unit": 50, "cost": 300, "price": 450},
    {"name": "Humic Acid", "category": "fertilizer", "brand": "Aries", "unit": "liter", "qty_per_unit": 1, "cost": 350, "price": 500},
]

MILL_PRODUCTS = ["Rice", "Wheat Flour", "Maize Flour", "Groundnut Oil", "Sunflower Oil", "Turmeric Powder", "Chilli Powder", "Bengal Gram Dal", "Lentil Dal"]

# ─── Helper functions ─────────────────────────────────────────────────────────

def rand_phone():
    prefixes = ["6", "7", "8", "9"]
    return random.choice(prefixes) + "".join([str(random.randint(0, 9)) for _ in range(9)])

def rand_email(name, idx, role):
    domain = random.choice(["gmail.com", "yahoo.co.in", "rediffmail.com", "outlook.com"])
    clean = name.lower().replace(" ", ".").replace("'", "")
    return f"{clean}.{role}{idx}@{domain}"

def rand_account():
    return "".join([str(random.randint(0, 9)) for _ in range(random.choice([11, 12, 13, 14]))])

def rand_ifsc(bank_code, state_idx):
    return f"{bank_code}0{state_idx:02d}{random.randint(100, 999)}"

def rand_aadhaar_last4():
    return "".join([str(random.randint(0, 9)) for _ in range(4)])

def rand_aadhaar_full():
    return "".join([str(random.randint(0, 9)) for _ in range(12)])

def rand_pan():
    letters = string.ascii_uppercase
    return "".join(random.choices(letters, k=5)) + "".join([str(random.randint(0, 9)) for _ in range(4)]) + random.choice(letters)

def rand_date_between(start, end):
    delta = (end - start).days
    if delta <= 0:
        return start
    return start + timedelta(days=random.randint(0, delta))

def rand_name(gender="male"):
    if gender == "female":
        first = random.choice(FEMALE_FIRST_NAMES)
    else:
        first = random.choice(MALE_FIRST_NAMES)
    last = random.choice(LAST_NAMES)
    return f"{first} {last}"

def rand_father():
    return f"{random.choice(FATHER_NAMES)} {random.choice(LAST_NAMES)}"

def rand_location():
    state_data = random.choice(LOCATIONS)
    state = state_data[0]
    dist_data = random.choice(state_data[1])
    district = dist_data[0]
    mandal = dist_data[1]
    village = random.choice(dist_data[2])
    pincode = dist_data[3] + str(random.randint(100, 999))
    house_no = f"{random.randint(1, 500)}/{random.choice(string.ascii_uppercase)}"
    street = random.choice(["Main Road", "Station Road", "Market Street", "Temple Road", "Gandhi Nagar",
                            "Nehru Colony", "Rajiv Nagar", "Indira Colony", "Ambedkar Street", "Bus Stand Road",
                            "Railway Colony", "Bypass Road", "College Road", "Hospital Road", "Bank Street"])
    return state, district, mandal, village, pincode, house_no, street

def rand_location_for_state(state_idx):
    state_data = LOCATIONS[state_idx]
    state = state_data[0]
    dist_data = random.choice(state_data[1])
    district = dist_data[0]
    mandal = dist_data[1]
    village = random.choice(dist_data[2])
    pincode = dist_data[3] + str(random.randint(100, 999))
    house_no = f"{random.randint(1, 500)}/{random.choice(string.ascii_uppercase)}"
    street = random.choice(["Main Road", "Station Road", "Market Street", "Temple Road", "Gandhi Nagar",
                            "Nehru Colony", "Rajiv Nagar", "Indira Colony", "Ambedkar Street", "Bus Stand Road"])
    return state, district, mandal, village, pincode, house_no, street

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

def style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

def auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 40)


# ═══════════════════════════════════════════════════════════════════════════════
#                          GENERATE EXCEL
# ═══════════════════════════════════════════════════════════════════════════════

def generate():
    wb = Workbook()
    print("🌾 AgriFlow Seed Data Generator")
    print("=" * 50)

    # ─── Sheet 1: Farmers ─────────────────────────────────────────────────────
    ws_farmers = wb.active
    ws_farmers.title = "farmers"
    farmer_headers = [
        "email", "phone_number", "password", "full_name", "farmer_id",
        "father_husband_name", "gender", "relation_type",
        "house_no", "street", "village", "mandal", "district", "state", "pincode",
        "total_area", "aadhaar_last_4", "bank_name", "account_number", "ifsc_code",
        "land_serial_1", "land_area_1", "land_serial_2", "land_area_2",
    ]
    ws_farmers.append(farmer_headers)

    print(f"  📝 Generating 10,000 farmers...")
    for i in range(1, 10001):
        state_idx = (i - 1) // 1000  # 1000 per state
        if state_idx >= len(LOCATIONS):
            state_idx = state_idx % len(LOCATIONS)
        state, district, mandal, village, pincode, house_no, street = rand_location_for_state(state_idx)
        
        gender = random.choice(["male"] * 8 + ["female"] * 2)  # 80% male
        name = rand_name(gender)
        father = rand_father()
        relation = "S/O" if gender == "male" else "W/O"
        total_area = round(random.uniform(1.0, 25.0), 1)
        bank_name, bank_code = random.choice(BANKS)
        
        # Land records (1 or 2)
        land1_serial = f"SY-{random.randint(100, 999)}/{random.randint(1, 50)}"
        land1_area = round(total_area * random.uniform(0.5, 0.8), 1)
        land2_serial = f"SY-{random.randint(100, 999)}/{random.randint(1, 50)}" if total_area > 3 else ""
        land2_area = round(total_area - land1_area, 1) if land2_serial else ""

        ws_farmers.append([
            rand_email(name, i, "farmer"),
            rand_phone(),
            "Farmer@123",
            name,
            f"FRM-{state[:2].upper()}-{i:05d}",
            father,
            gender,
            relation,
            house_no, street, village, mandal, district, state, pincode,
            total_area,
            rand_aadhaar_last4(),
            bank_name,
            rand_account(),
            rand_ifsc(bank_code, state_idx),
            land1_serial, land1_area, land2_serial, land2_area,
        ])

    style_header(ws_farmers)
    auto_width(ws_farmers)
    print(f"    ✅ 10,000 farmers generated")

    # ─── Sheet 2: Shops ───────────────────────────────────────────────────────
    ws_shops = wb.create_sheet("shops")
    shop_headers = [
        "email", "phone_number", "password", "full_name",
        "shop_name", "license_number", "shop_id", "father_name", "relation_type", "owner_name",
        "aadhaar_number", "pan_number",
        "shop_address", "landmark",
        "house_no", "street", "village", "mandal", "district", "state", "pincode",
        "bank_name", "account_number", "ifsc_code",
    ]
    ws_shops.append(shop_headers)

    shop_name_prefixes = ["Sri", "Sai", "Maa", "Jai", "Om", "Balaji", "Krishna", "Lakshmi", "Shiva", "Durga"]
    shop_name_suffixes = ["Agri Centre", "Seeds & Fertilizers", "Farm Store", "Kisan Seva Kendra", "Agro Agencies",
                          "Fertilizer Depot", "Agri Inputs", "Farm Solutions", "Pesticide Centre", "Seed House"]

    print(f"  🏪 Generating 500 shops...")
    for i in range(1, 501):
        state_idx = (i - 1) // 50  # 50 per state
        if state_idx >= len(LOCATIONS):
            state_idx = state_idx % len(LOCATIONS)
        state, district, mandal, village, pincode, house_no, street = rand_location_for_state(state_idx)
        
        name = rand_name("male")
        shop_name = f"{random.choice(shop_name_prefixes)} {random.choice(shop_name_suffixes)}"
        bank_name, bank_code = random.choice(BANKS)
        
        ws_shops.append([
            rand_email(name, i, "shop"),
            rand_phone(),
            "Shop@123",
            name,
            shop_name,
            f"LIC-{state[:3].upper()}-{random.randint(10000, 99999)}",
            f"SHP-{state[:2].upper()}-{i:04d}",
            rand_father(),
            "S/O",
            name,
            rand_aadhaar_full(),
            rand_pan(),
            f"{house_no}, {street}, {village}",
            random.choice(["Near Bus Stand", "Near Market Yard", "Main Road Junction", "Near Railway Station", "Near Temple"]),
            house_no, street, village, mandal, district, state, pincode,
            bank_name,
            rand_account(),
            rand_ifsc(bank_code, state_idx),
        ])

    style_header(ws_shops)
    auto_width(ws_shops)
    print(f"    ✅ 500 shops generated")

    # ─── Sheet 3: Mills ───────────────────────────────────────────────────────
    ws_mills = wb.create_sheet("mills")
    mill_headers = [
        "email", "phone_number", "password", "full_name",
        "mill_name", "license_number", "mill_id", "father_name", "relation_type", "owner_name",
        "aadhaar_number", "pan_number",
        "house_no", "street", "village", "mandal", "district", "state", "pincode",
        "location_text",
        "bank_name", "account_number", "ifsc_code",
    ]
    ws_mills.append(mill_headers)

    mill_name_prefixes = ["Sri", "Sai", "Royal", "Modern", "Lakshmi", "Ganesh", "Balaji", "Venkata", "National", "Star"]
    mill_name_suffixes = ["Rice Mill", "Flour Mill", "Oil Mill", "Processing Unit", "Industries",
                          "Agro Industries", "Rice & Flour Mill", "Dal Mill", "Spice Mill", "Rice Factory"]

    print(f"  🏭 Generating 200 mills...")
    for i in range(1, 201):
        state_idx = (i - 1) // 20  # 20 per state
        if state_idx >= len(LOCATIONS):
            state_idx = state_idx % len(LOCATIONS)
        state, district, mandal, village, pincode, house_no, street = rand_location_for_state(state_idx)
        
        name = rand_name("male")
        mill_name = f"{random.choice(mill_name_prefixes)} {random.choice(mill_name_suffixes)}"
        bank_name, bank_code = random.choice(BANKS)
        
        ws_mills.append([
            rand_email(name, i, "mill"),
            rand_phone(),
            "Mill@123",
            name,
            mill_name,
            f"MFG-{state[:3].upper()}-{random.randint(10000, 99999)}",
            f"MIL-{state[:2].upper()}-{i:03d}",
            rand_father(),
            "S/O",
            name,
            rand_aadhaar_full(),
            rand_pan(),
            house_no, street, village, mandal, district, state, pincode,
            f"Industrial Area, {village}, {district}",
            bank_name,
            rand_account(),
            rand_ifsc(bank_code, state_idx),
        ])

    style_header(ws_mills)
    auto_width(ws_mills)
    print(f"    ✅ 200 mills generated")

    # ─── Sheet 4: Customers ───────────────────────────────────────────────────
    ws_customers = wb.create_sheet("customers")
    customer_headers = [
        "email", "phone_number", "password", "full_name",
        "father_name", "relation_type", "id_number",
        "house_no", "street", "village", "mandal", "district", "state", "pincode",
        "bank_name", "account_number", "ifsc_code",
    ]
    ws_customers.append(customer_headers)

    print(f"  🛒 Generating 2,000 customers...")
    for i in range(1, 2001):
        state_idx = (i - 1) // 200  # 200 per state
        if state_idx >= len(LOCATIONS):
            state_idx = state_idx % len(LOCATIONS)
        state, district, mandal, village, pincode, house_no, street = rand_location_for_state(state_idx)
        
        gender = random.choice(["male", "female"])
        name = rand_name(gender)
        relation = "S/O" if gender == "male" else random.choice(["W/O", "D/O"])
        bank_name, bank_code = random.choice(BANKS)
        
        ws_customers.append([
            rand_email(name, i, "cust"),
            rand_phone(),
            "Customer@123",
            name,
            rand_father(),
            relation,
            rand_aadhaar_full(),
            house_no, street, village, mandal, district, state, pincode,
            bank_name,
            rand_account(),
            rand_ifsc(bank_code, state_idx),
        ])

    style_header(ws_customers)
    auto_width(ws_customers)
    print(f"    ✅ 2,000 customers generated")

    # ─── Save ─────────────────────────────────────────────────────────────────
    output_path = os.path.join(os.path.dirname(__file__), "agriflow_seed_data.xlsx")
    wb.save(output_path)
    print(f"\n{'=' * 50}")
    print(f"📦 Excel saved: {output_path}")
    print(f"   Sheets: farmers(10,000) | shops(500) | mills(200) | customers(2,000)")
    print(f"   Total: 12,700 users")
    return output_path


if __name__ == "__main__":
    generate()
